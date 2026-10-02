"""Unabhaengige Nachrechnung in Python (ohne Excel-Formeln).

Liest die Werte aus Bedarf_Bestand_LST (Bezirks-Blaetter) und LST_Zielzustand_gesamt und rechnet
alle Zahlen des Dashboards nach: Stufen je Zeitpunkt, Soll-Ist, Status je Mitarbeiter, Hinweise.
verify.py vergleicht das Ergebnis mit den von LibreOffice berechneten Dashboard-Zellen.

Regeln (siehe README):
  * Stufen nur aus L-P: Azubi (L) -> Arb LST (M) -> Wmech (N) -> SigMech (O) -> SigMech RBEG (P).
    Jeder zaehlt einmal auf der rechtesten erreichten Stufe. "x" in U = Wmech erreicht, "x" in V = SigMech erreicht.
    Heute: nur "x". Ende Jahr J: "x" oder Jahr <= J. Nach Plan: alle Jahre.
  * Teamleiter zusaetzlich ueber S ("x" bzw. Jahr <= J).
  * Status je Mitarbeiter an der Ziel-Qualifikation (G): Fertig / Abschluss <Jahr> / Ueberfaellig / Fehlt / Ziel unklar.
"""
from collections import Counter, OrderedDict
from datetime import datetime, timedelta

from openpyxl import load_workbook

STAGE_COLS = ["L", "M", "N", "O", "P"]
STAGE_NAMES = ["Azubi", "Arbeiter LST", "Wmech", "SigMech", "SigMech RBEG"]
CATS = ["Arbeiter LST", "Wmech", "SigMech", "SigMech RBEG", "Teamleiter"]
CHECK_COLS = ["L", "M", "N", "O", "P", "S", "U", "V"]
N_YEARS = 7
EMPTY = 9999
# Ziel-/Ist-Qualifikation (klein geschrieben) -> Spalte L..S
QUALI_COL = {
    "azubi": "L", "arb lst": "M", "arbeiter lst": "M",
    "weichmech": "N", "weichenmechaniker": "N", "wmech": "N",
    "sigmech": "O", "signalmechaniker": "O",
    "sigmech rbeg": "P", "signalmechaniker rbeg": "P",
    "kennziffer 4": "Q", "ihk-meister": "R", "teamleiter": "S", "tl": "S",
}
VP_OF = {"N": "U", "O": "V"}
STATUS_FIXED = {1: "Fertig", 7: "Überfällig", 8: "Fehlt (nichts geplant)", 9: "Ziel unklar"}


def status_label(code, bj):
    if code in STATUS_FIXED:
        return STATUS_FIXED[code]
    off = code - 2
    return f"Abschluss ab {bj + off}" if code == 6 else f"Abschluss {bj + off}"


def as_text(v):
    """Wie TRIM(CLEAN(Zelle&"")) in Excel."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "WAHR" if v else "FALSCH"
    if isinstance(v, datetime):
        v = (v - datetime(1899, 12, 30)).days
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = "".join(ch for ch in str(v) if ord(ch) >= 32)          # CLEAN
    return " ".join(part for part in s.split(" ") if part)    # TRIM


def parse(text):
    """0 = "x", Jahr, EMPTY = leer oder kein gueltiger Eintrag."""
    if text == "":
        return EMPTY
    if text.lower() == "x":
        return 0
    try:
        f = float(text.replace(",", "."))
    except ValueError:
        return EMPTY
    if f > 2100:
        return (datetime(1899, 12, 30) + timedelta(days=int(f))).year
    if f >= 1900:
        return int(f + 0.5)
    if 1 <= f <= 99:
        return 2000 + int(f + 0.5)
    return EMPTY


def stage_at(v, t):
    s = 0
    for i, c in enumerate(STAGE_COLS, start=1):
        if v[c] <= t:
            s = i
    if v["U"] == 0:
        s = max(s, 3)
    if v["V"] == 0:
        s = max(s, 4)
    return s


def compute(src_path, zz_path, districts, bj, davon_year=None, zz_sheet=None):
    """districts: Liste (Blattname, Name im Zielzustand). Gibt ein dict mit allen Ergebnissen zurueck.
    zz_sheet: Blatt mit dem Soll (Standard: erstes Blatt der Zielzustand-Datei)."""
    davon_year = bj + 1 if davon_year is None else davon_year
    wb = load_workbook(src_path, data_only=True)
    zz_wb = load_workbook(zz_path, data_only=True)
    zz = zz_wb[zz_sheet] if zz_sheet else zz_wb.worksheets[0]
    zz_col = {as_text(c.value): c.column for c in zz[1] if c.value}
    zz_row = {as_text(zz.cell(r, 1).value): r for r in range(1, zz.max_row + 1)}
    thresholds = [0] + [bj + k for k in range(N_YEARS)] + [EMPTY - 1]   # heute, Jahre, nach Plan
    people, out = [], OrderedDict()
    for sheet, zz_name in districts:
        ws = wb[sheet]
        for r in range(2, 102):
            name = as_text(ws[f"B{r}"].value)
            if name == "" or name == "Name":
                continue
            raw = {c: as_text(ws[f"{c}{r}"].value) for c in CHECK_COLS + ["Q", "R"]}
            p = dict(district=sheet, row=r, name=name, first=as_text(ws[f"C{r}"].value),
                     ziel=as_text(ws[f"G{r}"].value), ist=as_text(ws[f"J{r}"].value), raw=raw,
                     v={c: parse(raw[c]) for c in CHECK_COLS + ["Q", "R"]})
            p["stage"] = [stage_at(p["v"], t) for t in thresholds]
            p["tl"] = [1 if p["v"]["S"] <= t else 0 for t in thresholds]
            people.append(p)

    dup = Counter((p["name"].lower(), p["first"].lower()) for p in people)
    for p in people:
        v, raw = p["v"], p["raw"]
        col = QUALI_COL.get(p["ziel"].lower(), "") if p["ziel"] else ""
        if col == "":
            code, year = 9, None
        else:
            y = v[col]
            vp_x = col in VP_OF and v[VP_OF[col]] == 0
            year = y if 0 < y < EMPTY else None
            if y == 0 or vp_x:
                code = 1
            elif year is None:
                code = 8
            elif year < bj:
                code = 7
            else:
                code = min(6, 2 + year - bj)
        p.update(status=code, status_text=status_label(code, bj), year=year, ziel_col=col)

        # Hinweise
        h = []
        past = [f"{c} {v[c]}" for c in CHECK_COLS if 0 < v[c] < EMPTY and v[c] < bj]
        h.append("Jahr vorbei ohne „x“: " + ", ".join(past) if past else "")
        bad = [f"{c} „{raw[c]}“" for c in CHECK_COLS if raw[c] != "" and v[c] == EMPTY]
        h.append("Weder „x“ noch Jahr: " + ", ".join(bad) if bad else "")
        h.append(f"Ziel „{p['ziel']}“: nichts geplant in Spalte {col}" + (f"/{VP_OF[col]}" if col in VP_OF else "")
                 if code == 8 else "")
        h.append(("Keine Ziel-Qualifikation eingetragen" if p["ziel"] == "" else f"Ziel „{p['ziel']}“ hat keine eigene Spalte")
                 if code == 9 else "")
        h.append("Heute auf keiner Stufe (kein „x“ in L–P)" if p["stage"][0] == 0 else "")
        ist = p["ist"].lower()
        ist_nr = 1 if ist[:5] == "azubi" else ("LMNOPQRS".index(QUALI_COL[ist]) + 1 if ist in QUALI_COL else 0)
        if 1 <= ist_nr <= 5 and ist_nr != p["stage"][0]:
            now = STAGE_NAMES[p["stage"][0] - 1] if p["stage"][0] else "keine Stufe"
            h.append(f"Ist-Qualifikation „{p['ist']}“, laut L–P heute {now}")
        elif ist_nr == 8 and v["S"] != 0:
            h.append(f"Ist-Qualifikation „{p['ist']}“, aber kein „x“ in S")
        else:
            h.append("")
        gap = []
        if raw["U"] != "" and raw["N"] == "":
            gap.append("U eingetragen, aber N leer")
        if raw["V"] != "" and raw["O"] == "":
            gap.append("V eingetragen, aber O leer")
        if raw["P"] != "" and raw["O"] == "" and v["V"] != 0:
            gap.append("P eingetragen, aber O leer")
        h.append("Stufe fehlt: " + "; ".join(gap) if gap else "")
        n = dup[(p["name"].lower(), p["first"].lower())]
        h.append(f"Name und Vorname stehen {n}× in der Liste" if n > 1 else "")
        p["flags"] = h
        p["hints"] = " · ".join(x for x in h if x)

    for sheet, zz_name in districts:
        ps = [p for p in people if p["district"] == sheet]
        c = zz_col.get(zz_name)
        soll = {}
        for cat in CATS:
            val = zz.cell(zz_row[cat], c + 3).value if c and cat in zz_row else 0
            soll[cat] = val if isinstance(val, (int, float)) else 0
        counts = {}
        for k in range(len(thresholds)):
            row = {name: sum(1 for p in ps if p["stage"][k] == i) for i, name in enumerate(STAGE_NAMES, start=1)}
            row["Teamleiter"] = sum(p["tl"][k] for p in ps)
            row["ohne Stufe"] = sum(1 for p in ps if p["stage"][k] == 0)
            counts[k] = row
        compare = {cat: [compare_text(counts[k][cat], soll[cat]) for k in range(len(thresholds))] for cat in CATS}
        missing = [sum(max(0, soll[cat] - counts[k][cat]) for cat in CATS) for k in range(len(thresholds))]
        more = [sum(max(0, counts[k][cat] - soll[cat]) for cat in CATS) for k in range(len(thresholds))]
        st = Counter(p["status"] for p in ps)
        out[sheet] = dict(
            people=sorted(ps, key=lambda p: (p["status"], p["row"])), soll=soll, counts=counts, compare=compare,
            missing=missing, more=more, status=st, headcount=len(ps),
            in_training=sum(st[i] for i in range(2, 7)),
            davon=sum(1 for p in ps if 2 <= p["status"] <= 6 and p["year"] == davon_year),
            by_ziel=Counter((ziel_group(p), p["status"]) for p in ps))
    return dict(districts=out, people=people, thresholds=thresholds, bj=bj, davon_year=davon_year)


def ziel_group(p):
    col = p["ziel_col"]
    return {"M": "Arbeiter LST", "N": "Wmech", "O": "SigMech", "P": "SigMech RBEG", "S": "Teamleiter"}.get(col, "Sonstige")


def compare_text(ist, soll):
    if ist == soll:
        return "Soll erreicht"
    if ist < soll:
        return "es fehlt 1" if soll - ist == 1 else f"es fehlen {soll - ist}"
    return f"{ist - soll} mehr als Soll"


if __name__ == "__main__":
    import sys
    from build_dashboard import DISTRICTS
    res = compute(sys.argv[1], sys.argv[2], [(a, b) for a, b, _ in DISTRICTS], int(sys.argv[3]) if len(sys.argv) > 3 else datetime.now().year)
    for sheet, d in res["districts"].items():
        print(f"== {sheet}: {d['headcount']} Mitarbeiter, Soll {d['soll']}")
        for name in STAGE_NAMES + ["Teamleiter", "ohne Stufe"]:
            print(f"   {name:14s}", [d["counts"][k][name] for k in range(len(res['thresholds']))])
