"""Vergleicht die von LibreOffice berechneten Dashboard-Werte mit der unabhaengigen Nachrechnung (reference.py).

    python verify.py <berechnete Datei (KEEP_CALC)> <Bedarf_Bestand_LST.xlsx> <LST_Zielzustand_gesamt.xlsx> [Bezugsjahr]

Geprueft werden: Auswertung je Mitarbeiter (Blatt Daten), alle Zahlentabellen und Kennzahlen der Bezirksseiten,
Mitarbeiterlisten, Uebersicht, Soll-Ist und Pruefliste. Mit --liste wird zusaetzlich die Liste der Hinweise ausgegeben.
"""
import sys
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import build_sheets as B
from build_dashboard import DISTRICTS
from reference import CATS, STAGE_NAMES, compute


class Check:
    def __init__(self):
        self.n = 0
        self.bad = []

    def eq(self, where, got, want):
        self.n += 1
        if isinstance(got, float) and got.is_integer():
            got = int(got)
        if got is None:
            got = ""
        if got != want:
            self.bad.append(f"{where}: Dashboard {got!r} ≠ Python {want!r}")


def verify(calc_path, src, zz, bj, zz_sheet=None):
    districts = [(a, b) for a, b, _ in DISTRICTS]
    ref = compute(src, zz, districts, bj, zz_sheet=zz_sheet)
    wb = load_workbook(calc_path, data_only=True)
    ck = Check()
    B.DATA_LAST = 1 + len(districts) * B.SRC_ROWS

    # Blatt Daten: je Mitarbeiter
    ws = wb["Daten"]
    col = {k: B.DC[k] for k in B.DC}
    by_pos = {}
    for r in range(2, B.DATA_LAST + 1):
        if ws[f"{col['ma']}{r}"].value == 1:
            by_pos[(ws[f"{col['bez']}{r}"].value, ws[f"{col['row']}{r}"].value)] = r
    ck.eq("Daten: Anzahl Mitarbeiter", len(by_pos), len(ref["people"]))
    for p in ref["people"]:
        r = by_pos.get((p["district"], p["row"]))
        where = f"Daten {p['district']} Zeile {p['row']} ({p['name']})"
        if r is None:
            ck.bad.append(f"{where}: fehlt im Blatt Daten")
            continue
        for k in range(B.N_TL):
            ck.eq(f"{where} Stufe[{k}]", ws[f"{col[f'st{k}']}{r}"].value, p["stage"][k])
            ck.eq(f"{where} Teamleiter[{k}]", ws[f"{col[f'tl{k}']}{r}"].value, p["tl"][k])
        ck.eq(f"{where} Status", ws[f"{col['status']}{r}"].value, p["status"])
        ck.eq(f"{where} Status-Text", ws[f"{col['stext']}{r}"].value, p["status_text"])
        for i in range(8):
            ck.eq(f"{where} Hinweis {i + 1}", ws[f"{col[f'f{i + 1}']}{r}"].value, p["flags"][i])

    # Bezirksseiten
    for sheet, _ in districts:
        d = ref["districts"][sheet]
        ws = wb[B.dash_name(sheet)]
        w = f"Dashboard {sheet}"
        for name in STAGE_NAMES + ["Teamleiter", "ohne Stufe"]:
            for k, c in enumerate(B.TBL):
                ck.eq(f"{w} {name} {c}", ws[f"{c}{B.DT_ROWS[name]}"].value, d["counts"][k][name])
        for k, c in enumerate(B.TBL):
            ck.eq(f"{w} gesamt {c}", ws[f"{c}{B.DT_ROWS['gesamt']}"].value, d["headcount"])
            ck.eq(f"{w} fehlende Stellen {c}", ws[f"{c}{B.DT_MISSING}"].value, d["missing"][k])
            ck.eq(f"{w} mehr als Soll {c}", ws[f"{c}{B.DT_MORE}"].value, d["more"][k])
        for cat in CATS:
            ck.eq(f"{w} Soll {cat}", ws[f"{B.T_SOLL}{B.DT_ROWS[cat]}"].value, d["soll"][cat])
            for k, c in enumerate(B.TBL):
                ck.eq(f"{w} Abgleich {cat} {c}", ws[f"{c}{B.DT_CMP[cat]}"].value, d["compare"][cat][k])
        for g, r in B.DT_ST.items():
            for code in range(1, 10):
                c = get_column_letter(B.STAT_COL1 + code - 1)
                ck.eq(f"{w} Status {g}/{code}", ws[f"{c}{r}"].value, d["by_ziel"].get((g, code), 0))
        for code in range(1, 10):
            c = get_column_letter(B.STAT_COL1 + code - 1)
            ck.eq(f"{w} Status Summe {code}", ws[f"{c}{B.DT_ST_SUM}"].value, d["status"].get(code, 0))
        ck.eq(f"{w} KPI Mitarbeiter", ws["B6"].value, d["headcount"])
        ck.eq(f"{w} KPI Fertig", ws["E6"].value, d["status"].get(1, 0))
        ck.eq(f"{w} KPI In Ausbildung", ws["H6"].value, d["in_training"])
        ck.eq(f"{w} KPI davon", (ws["H7"].value or "").split(" mit Abschluss")[0], f"davon {d['davon']}")
        ck.eq(f"{w} KPI Fehlt", ws["K6"].value, d["status"].get(8, 0))
        ck.eq(f"{w} KPI fehlende Stellen", ws["N6"].value, d["missing"][0])
        for i, p in enumerate(d["people"][:B.LIST_ROWS]):
            r = B.DR_LIST_FIRST + i
            ck.eq(f"{w} Liste {i + 1}", (ws[f"C{r}"].value, ws[f"E{r}"].value, ws[f"N{r}"].value),
                  (p["name"], p["first"], p["status_text"]))
            now = (STAGE_NAMES[p["stage"][0] - 1] if p["stage"][0] else "–") + (" + Teamleiter" if p["tl"][0] else "")
            ck.eq(f"{w} Liste {i + 1} Stufe heute", ws[f"K{r}"].value, now)
            ck.eq(f"{w} Liste {i + 1} Abschluss", ws[f"M{r}"].value, p["year"] if 2 <= p["status"] <= 7 else "–")
        ck.eq(f"{w} Liste Ende", ws[f"C{B.DR_LIST_FIRST + len(d['people'])}"].value, "")

    # Uebersicht
    ws = wb["Übersicht"]
    for i, (sheet, _) in enumerate(districts):
        d = ref["districts"][sheet]
        r = B.OV_FIRST + i
        for code in range(1, 10):
            ck.eq(f"Übersicht {sheet} Status {code}", ws.cell(r, 4 + code).value, d["status"].get(code, 0))
        ck.eq(f"Übersicht {sheet} Hinweise", ws[f"Q{r}"].value, sum(1 for p in d["people"] if p["hints"]))
    tot = lambda key: sum(d[key] for d in ref["districts"].values())
    ck.eq("Übersicht KPI Mitarbeiter", ws["B6"].value, tot("headcount"))
    ck.eq("Übersicht KPI In Ausbildung", ws["H6"].value, tot("in_training"))
    ck.eq("Übersicht KPI davon", (ws["H7"].value or "").split(" mit Abschluss")[0], f"davon {tot('davon')}")
    ck.eq("Übersicht KPI fehlende Stellen", ws["N6"].value, sum(d["missing"][0] for d in ref["districts"].values()))

    # Soll-Ist
    ws = wb["Soll-Ist"]
    for name in STAGE_NAMES + ["Teamleiter", "ohne Stufe"]:
        for k, c in enumerate(B.TBL):
            ck.eq(f"Soll-Ist gesamt {name} {c}", ws[f"{c}{B.SI_TG_ROWS[name]}"].value,
                  sum(d["counts"][k][name] for d in ref["districts"].values()))
    for k, c in enumerate(B.TBL):
        ck.eq(f"Soll-Ist fehlende Stellen gesamt {c}", ws[f"{c}{B.SI_TG_MISS}"].value,
              sum(d["missing"][k] for d in ref["districts"].values()))
        ck.eq(f"Soll-Ist mehr als Soll gesamt {c}", ws[f"{c}{B.SI_TG_MORE}"].value,
              sum(d["more"][k] for d in ref["districts"].values()))
    for i, (sheet, _) in enumerate(districts):
        d = ref["districts"][sheet]
        r = B.SI_TB_FIRST + i
        ck.eq(f"Soll-Ist {sheet} Soll", ws[f"{B.T_SOLL}{r}"].value, sum(d["soll"].values()))
        for k, c in enumerate(B.TBL):
            ck.eq(f"Soll-Ist {sheet} fehlende Stellen {c}", ws[f"{c}{r}"].value, d["missing"][k])
    for cat in CATS:
        ck.eq(f"Soll-Ist Soll {cat}", ws[f"{B.T_SOLL}{B.SI_TG_ROWS[cat]}"].value,
              sum(d["soll"][cat] for d in ref["districts"].values()))

    # Pruefliste
    ws = wb["Prüfliste"]
    flagged = [p for p in ref["people"] if p["hints"]]
    for i, p in enumerate(flagged[:B.PRUEF_ROWS]):
        r = B.PR_FIRST + i
        ck.eq(f"Prüfliste {i + 1}", (ws[f"C{r}"].value, ws[f"D{r}"].value, ws[f"E{r}"].value, ws[f"H{r}"].value),
              (p["district"], p["name"], p["first"], p["hints"]))
    ck.eq("Prüfliste Ende", ws[f"D{B.PR_FIRST + len(flagged)}"].value, "")
    return ck, ref


def anomaly_list(ref):
    lines = []
    for p in ref["people"]:
        if p["hints"]:
            lines.append(f"| {p['district']} | {p['name']}, {p['first']} | {p['ziel'] or '–'} | {p['hints']} |")
    return lines


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    bj = int(args[3]) if len(args) > 3 else date.today().year
    ck, ref = verify(args[0], args[1], args[2], bj)
    print(f"{ck.n} Werte verglichen, {len(ck.bad)} Abweichungen")
    for b in ck.bad[:60]:
        print("  ", b)
    if "--liste" in sys.argv:
        print("\n".join(anomaly_list(ref)))
    sys.exit(1 if ck.bad else 0)
