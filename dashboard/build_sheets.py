"""Erzeugt die Dashboard-Blaetter als eigenstaendige Arbeitsmappe:

    Uebersicht, Soll-Ist, Dashboard <Bezirk> (je Bezirk), Pruefliste, Zielzustand, Einstellungen, Daten

Die Formeln verweisen auf die Bezirks-Blaetter der Original-Datei (Bedarf_Bestand_LST);
merge_into_original.py setzt die Blaetter anschliessend dort ein. Das Blatt Zielzustand
ist eine Kopie aus LST_Zielzustand_gesamt.xlsx und dient als Soll.

Zaehlregel Soll-Ist / Verlauf: nur Spalten L-P (Azubi -> Arb LST -> Wmech -> SigMech -> SigMech RBEG).
Jeder Mitarbeiter zaehlt einmal auf der rechtesten erreichten Stufe ("x" in U = Wmech, "x" in V = SigMech).
Heute = nur "x", Ende eines Jahres = "x" oder Jahr <= diesem Jahr, nach Plan = alle Jahre.
Teamleiter zaehlen zusaetzlich ueber Spalte S. Ziel- (G) und Ist-Qualifikation (J) zaehlen hier nicht.
Status je Mitarbeiter: gemessen an der Ziel-Qualifikation (G).
"""
from copy import copy
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint, Series, SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter, quote_sheetname
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties


@dataclass
class District:
    sheet: str      # Blattname in Bedarf_Bestand_LST
    zz_name: str    # Name des Bezirks im Zielzustand (Zeile 1)


def dash_name(sheet):
    return f"Dashboard {sheet}"


# ---------------------------------------------------------------- Konstanten
SRC_FIRST_ROW = 2          # erste gelesene Zeile je Bezirks-Blatt (Kopfzeilen werden ueber "Name" erkannt)
SRC_ROWS = 100             # gelesene Zeilen je Bezirks-Blatt (2..101)
SRC_MAXROW = 250           # Ende der begrenzten Bereiche in den Formeln (keine ganzen Spalten)
LIST_ROWS = 30             # Mitarbeiterliste je Bezirksseite
PRUEF_ROWS = 80            # Zeilen der Pruefliste
N_YEARS = 7                # Verlauf: Bezugsjahr .. Bezugsjahr+6
N_TL = N_YEARS + 2         # Zeitpunkte: heute, 7 Jahre, nach Plan

SRC_NAME, SRC_FIRST, SRC_ZIEL, SRC_IST = "B", "C", "G", "J"
SRC_ENTRY = ["L", "M", "N", "O", "P", "Q", "R", "S", "U", "V"]   # gelesene Eintraege (L..S zusammenhaengend)
CHECK_COLS = ["L", "M", "N", "O", "P", "S", "U", "V"]            # fuer Hinweise geprueft
STAGES = [("Azubi", "L"), ("Arbeiter LST", "M"), ("Wmech", "N"), ("SigMech", "O"), ("SigMech RBEG", "P")]
STAGE_NAMES = [s for s, _ in STAGES]
TL_NAME = "Teamleiter"
NO_STAGE = "ohne Stufe"
CATS = ["Arbeiter LST", "Wmech", "SigMech", "SigMech RBEG", "Teamleiter"]   # Zeilen im Zielzustand (Soll)
OTHER = "Sonstige"
GROUPS = CATS + [OTHER]
# Erwartete Ueberschriften (Kopfzeile 1 oder 2) - werden im Blatt Einstellungen per Formel geprueft
HEADERS = {"B": "Name", "C": "Vorname", "G": "Ziel-Qualifikation", "J": "Ist-Qualifikation", "L": "Azubi",
           "M": "Arb LST", "N": "Weichenmechaniker", "O": "Signalmechaniker", "P": "Signalmechaniker RBEG",
           "S": "Teamleiter"}
HEADER_PREFIX = {"U": "örtl", "V": "örtl"}

FONT = "Aptos Narrow"       # Schrift der Original-Datei
INK = "0B0B0B"
INK2 = "52514E"
MUTED = "898781"
HAIR = "E1E0D9"
TILE_BG = "F4F4F2"
HEAD_BG = "EDEDEA"
INPUT_BG = "FFF2CC"
LINK = "1C5CAB"
GOOD_BG, GOOD_INK = "E2F3E2", "0A6B0A"
BAD_BG, BAD_INK = "F9E0E0", "A12A2A"
MORE_BG, MORE_INK = "FDF0D5", "7A5300"

# Farben Soll-Ist
C_SOLL = INK2
C_HEUTE = "86B6EF"
C_PLAN = "2A78D6"

# Status-Nr -> (Farbe, Textfarbe auf der Farbe). Bezeichnungen stehen im Blatt Einstellungen.
STATUS = [
    (1, "0CA30C", INK),       # Fertig
    (2, "104281", "FFFFFF"),  # Abschluss Bezugsjahr
    (3, "1C5CAB", "FFFFFF"),  # Abschluss Bezugsjahr+1
    (4, "2A78D6", "FFFFFF"),  # +2
    (5, "5598E7", INK),       # +3
    (6, "86B6EF", INK),       # ab +4
    (7, "EC835A", INK),       # Ueberfaellig
    (8, "D03B3B", "FFFFFF"),  # Fehlt
    (9, "B5B3AC", INK),       # Ziel unklar
]
STATUS_TEXT = {
    1: ("Fertig", "„x“ in der Spalte der Ziel-Qualifikation (L–S). Bei Wmech und SigMech reicht auch ein „x“ bei der "
                  "örtlichen Verwendungsprüfung (U bzw. V)."),
    2: (None, "In Ausbildung: In der Spalte der Ziel-Qualifikation steht als Jahr das Bezugsjahr."),
    3: (None, "In Ausbildung: geplanter Abschluss im Bezugsjahr + 1."),
    4: (None, "In Ausbildung: geplanter Abschluss im Bezugsjahr + 2."),
    5: (None, "In Ausbildung: geplanter Abschluss im Bezugsjahr + 3."),
    6: (None, "In Ausbildung: geplanter Abschluss im Bezugsjahr + 4 oder später."),
    7: ("Überfällig", "Das geplante Jahr ist vorbei, aber es steht noch kein „x“. Bitte prüfen."),
    8: ("Fehlt (nichts geplant)", "In der Spalte der Ziel-Qualifikation steht weder „x“ noch ein Jahr."),
    9: ("Ziel unklar", "Keine Ziel-Qualifikation eingetragen, oder das Ziel hat keine eigene Spalte (z. B. Senior Expert LST)."),
}

# Zuordnung Text in Ziel-/Ist-Spalte -> Spalte L..S, oertl. Verwendungspruefung U/V, Kategorie (Zielzustand)
MAPPING = [
    ("Azubi", "L", "", OTHER),
    ("Arb LST", "M", "", "Arbeiter LST"),
    ("Arbeiter LST", "M", "", "Arbeiter LST"),
    ("Weichmech", "N", "U", "Wmech"),
    ("Weichenmechaniker", "N", "U", "Wmech"),
    ("Wmech", "N", "U", "Wmech"),
    ("Sigmech", "O", "V", "SigMech"),
    ("Signalmechaniker", "O", "V", "SigMech"),
    ("Sigmech RBEG", "P", "", "SigMech RBEG"),
    ("Signalmechaniker RBEG", "P", "", "SigMech RBEG"),
    ("Kennziffer 4", "Q", "", OTHER),
    ("IHK-Meister", "R", "", OTHER),
    ("Teamleiter", "S", "", "Teamleiter"),
    ("TL", "S", "", "Teamleiter"),
    ("Senior Expert LST", "", "", OTHER),
    ("Umschüler EBET", "", "", OTHER),
    ("Umschüler", "", "", OTHER),
    ("Quereinsteiger", "", "", OTHER),
]

# Stufen-Farben (Verlauf): hell = frueh im Ausbildungsweg, dunkel = spaet; Teamleiter eigene Farbe
STAGE_COLORS = {NO_STAGE: ("E1E0D9", INK2), "Azubi": ("86B6EF", INK), "Arbeiter LST": ("5598E7", INK),
                "Wmech": ("2A78D6", "FFFFFF"), "SigMech": ("1C5CAB", "FFFFFF"), "SigMech RBEG": ("104281", "FFFFFF"),
                TL_NAME: ("EB6834", INK)}
STACK_ORDER = [NO_STAGE] + STAGE_NAMES + [TL_NAME]      # von unten nach oben

TXT_STATUS = (
    "So wird der Status gezählt: Maßstab ist die Ziel-Qualifikation (Spalte G). Fertig = „x“ in der Spalte dieses Ziels "
    "(bei Wmech/SigMech reicht auch „x“ bei der örtlichen Verwendungsprüfung U/V). Abschluss + Jahr = dort steht ein geplantes "
    "Jahr (z. B. 27 = 2027). Überfällig = das Jahr ist vorbei, aber noch kein „x“. Fehlt = dort steht nichts. "
    "Ziel unklar = kein Ziel eingetragen oder das Ziel hat keine eigene Spalte.")
TXT_COUNT = (
    "So wird je Stufe gezählt: nur mit den Spalten L–P (Azubi → Arb LST → Wmech → SigMech → SigMech RBEG). Jeder Mitarbeiter "
    "zählt genau einmal – auf der rechtesten Stufe, die erreicht ist. Heute = rechteste Spalte mit „x“ („x“ in U zählt als Wmech, "
    "„x“ in V als SigMech). Ende eines Jahres = rechteste Spalte mit „x“ oder mit einem Jahr bis zu diesem Jahr. Nach Plan = alle "
    "eingetragenen Jahre. Teamleiter zählen zusätzlich über Spalte S – ein Teamleiter mit „x“ in P zählt also auch bei "
    "SigMech RBEG. Soll = Spalte „Zielzustand“ im Blatt Zielzustand.")

# Einstellungen: feste Zellpositionen
E_YEAR = "Einstellungen!$C$4"
E_DAVON = "Einstellungen!$C$5"
E_STAT_FIRST, E_STAT_LAST = 9, 17
E_STAGE_FIRST, E_STAGE_LAST = 21, 25
E_BEZ_HDR, E_BEZ_FIRST, E_BEZ_LAST = 30, 31, 50
E_MAP_HDR, E_MAP_FIRST, E_MAP_LAST = 55, 56, 95
E_STAT_LABEL = f"Einstellungen!$C${E_STAT_FIRST}:$C${E_STAT_LAST}"
E_STAGE_NAMES = f"Einstellungen!$C${E_STAGE_FIRST}:$C${E_STAGE_LAST}"
E_BEZ_SHEET = f"Einstellungen!$B${E_BEZ_FIRST}:$B${E_BEZ_LAST}"
E_BEZ_ZZNAME = f"Einstellungen!$C${E_BEZ_FIRST}:$C${E_BEZ_LAST}"
E_BEZ_ZZCOL = f"Einstellungen!$D${E_BEZ_FIRST}:$D${E_BEZ_LAST}"
E_MAP_TEXT = f"Einstellungen!$B${E_MAP_FIRST}:$B${E_MAP_LAST}"
E_MAP_COLNO = f"Einstellungen!$D${E_MAP_FIRST}:$D${E_MAP_LAST}"
E_MAP_VPNO = f"Einstellungen!$F${E_MAP_FIRST}:$F${E_MAP_LAST}"
E_MAP_CAT = f"Einstellungen!$G${E_MAP_FIRST}:$G${E_MAP_LAST}"
ZZ_AREA = "Zielzustand!$A$1:$BZ$60"
ZZ_LABELS = "Zielzustand!$A$1:$A$60"


def stat_label_ref(code):
    return f"Einstellungen!$C${E_STAT_FIRST + code - 1}"


def year_ref(k):
    """Zeitpunkt k: 0 = heute, 1..7 = Ende Bezugsjahr+k-1, 8 = nach Plan."""
    return f"({E_YEAR}+{k - 1})"


def timeline_label(k):
    if k == 0:
        return "heute"
    if k == N_TL - 1:
        return "nach Plan"
    return f'=""&{year_ref(k)}'


# ---------------------------------------------------------------- Stil-Helfer
def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


HAIR_SIDE = Side(style="thin", color=HAIR)
BOTTOM_HAIR = Border(bottom=HAIR_SIDE)
TOP_INK = Border(top=Side(style="thin", color=INK2))
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
NF_COUNT = '0;-0;"·"'


def put(ws, ref, value, f=None, fl=None, al=None, nf=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fl:
        c.fill = fl
    if al:
        c.alignment = al
    if nf:
        c.number_format = nf
    if border:
        c.border = border
    return c


def merge_put(ws, rng_, value, **kw):
    first = rng_.split(":")[0]
    c = put(ws, first, value, **kw)
    ws.merge_cells(rng_)
    if kw.get("fl") or kw.get("border"):
        for row in ws[rng_]:
            for cell in row:
                if kw.get("fl"):
                    cell.fill = kw["fl"]
                if kw.get("border"):
                    cell.border = kw["border"]
    return c


def head(ws, ref, text, wrap=True, align="center"):
    al = (CENTER_WRAP if wrap else CENTER) if align == "center" else LEFT_WRAP
    if ":" in ref:
        return merge_put(ws, ref, text, f=font(9, True, INK2), fl=fill(HEAD_BG), al=al)
    return put(ws, ref, text, f=font(9, True, INK2), fl=fill(HEAD_BG), al=al)


def page_setup(ws, print_area, breaks=(), title_rows=None):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = print_area
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.45
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.oddFooter.right.text = "Seite &P von &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.left.size = 8
    if title_rows:
        ws.print_title_rows = title_rows
    if breaks:
        ws.row_breaks = RowBreak()
        for r in breaks:
            ws.row_breaks.append(Break(id=r))


# ---------------------------------------------------------------- Diagramm-Helfer
def text_props(size=9, color=INK2, bold=False):
    cp = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color,
                             latin=DFont(typeface=FONT), cs=DFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text, size=10):
    cp = CharacterProperties(sz=int(size * 100), b=True, solidFill=INK, latin=DFont(typeface=FONT), cs=DFont(typeface=FONT))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def no_line():
    return LineProperties(noFill=True)


def data_labels(color, pos=None, numfmt=None, bold=False, size=9):
    dl = DataLabelList()
    dl.showVal = True
    dl.showLegendKey = False
    dl.showCatName = False
    dl.showSerName = False
    dl.showPercent = False
    dl.showBubbleSize = False
    if pos:
        dl.position = pos
    if numfmt:
        dl.numFmt = numfmt
    dl.txPr = text_props(size, color, bold)
    return dl


def style_chart_frame(chart):
    chart.graphical_properties = GraphicalProperties(ln=no_line())
    chart.roundedCorners = False
    chart.visible_cells_only = False     # Hilfswerte stehen in ausgeblendeten Spalten


def cat_axis(ax, reverse=False, size=9):
    if reverse:
        ax.scaling.orientation = "maxMin"
    ax.delete = False
    ax.majorTickMark = "none"
    ax.txPr = text_props(size, INK2)
    ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="C3C2B7"))


def hide_value_axis(ax):
    ax.delete = True
    ax.majorGridlines = None
    ax.scaling.min = 0


def show_value_axis(ax, ymax=None, major=None, size=8):
    ax.delete = False
    ax.scaling.min = 0
    if ymax:
        ax.scaling.max = ymax
    if major:
        ax.majorUnit = major
    ax.numFmt = "0"
    ax.majorTickMark = "none"
    ax.txPr = text_props(size, MUTED)
    ax.graphicalProperties = GraphicalProperties(ln=no_line())
    ax.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=HAIR, w=6350)))


def legend_bottom(chart, size=9):
    chart.legend = Legend()
    chart.legend.position = "b"
    chart.legend.txPr = text_props(size, INK2)


def ref_series(idx, values_ref, title_ref, cats_ref):
    """Serie aus Bereichs-Texten ('Blatt'!$A$1:$B$1): Werte, Titel-Zelle, Kategorien."""
    s = Series(idx=idx, order=idx)
    s.val = NumDataSource(numRef=NumRef(f=values_ref))
    s.tx = SeriesLabel(strRef=StrRef(f=title_ref))
    s.cat = AxDataSource(strRef=StrRef(f=cats_ref))
    return s


def fill_series(s, color):
    s.graphicalProperties = GraphicalProperties(solidFill=color, ln=no_line())
    s.invertIfNegative = False
    return s


def rng(ws_title, c1, r1, c2=None, r2=None):
    return f"{quote_sheetname(ws_title)}!${c1}${r1}:${c2 or c1}${r2 or r1}"


def status_bar_chart(ws, val_row, val_c1, val_c2, cat_row):
    """Ein Balken je Status, jeweils in der Status-Farbe."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.style = 2
    ch.add_data(Reference(ws, min_col=val_c1, max_col=val_c2, min_row=val_row, max_row=val_row),
                from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=val_c1, max_col=val_c2, min_row=cat_row, max_row=cat_row))
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=STATUS[0][1], ln=no_line())
    for i, (_, color, _) in enumerate(STATUS):
        pt = DataPoint(idx=i, invertIfNegative=False)
        pt.graphicalProperties = GraphicalProperties(solidFill=color, ln=no_line())
        s.dPt.append(pt)
    s.invertIfNegative = False
    s.dLbls = data_labels(INK, pos="outEnd", numfmt="0;-0;0", bold=True)
    ch.gapWidth = 35
    ch.legend = None
    cat_axis(ch.x_axis, reverse=True)
    hide_value_axis(ch.y_axis)
    style_chart_frame(ch)
    return ch


def stacked_status_chart(ws, hdr_row, first_row, last_row, cat_col, c1, c2):
    """Gestapelte Balken: je Zeile die Anzahl je Status (Spalte)."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.style = 2
    ch.add_data(Reference(ws, min_col=c1, max_col=c2, min_row=hdr_row, max_row=last_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=first_row, max_row=last_row))
    for s, (_, color, txt) in zip(ch.series, STATUS):
        s.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill="FFFFFF", w=12700))
        s.invertIfNegative = False
        s.dLbls = data_labels(txt, pos="ctr", numfmt="0;-0;;")   # Nullen ausblenden
    ch.gapWidth = 45
    legend_bottom(ch)
    cat_axis(ch.x_axis, reverse=True)
    hide_value_axis(ch.y_axis)
    style_chart_frame(ch)
    return ch


def grouped_chart(cats_ref, series, horizontal=False, label_size=9):
    """Gruppierte Saeulen/Balken. series: Liste (Werte-Bereich, Titel-Zelle, Farbe)."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.style = 2
    for i, (values_ref, title_ref, color) in enumerate(series):
        s = fill_series(ref_series(i, values_ref, title_ref, cats_ref), color)
        s.dLbls = data_labels(INK2, pos="outEnd", numfmt="0;-0;0", size=label_size)
        ch.series.append(s)
    ch.gapWidth = 60
    ch.overlap = -8
    legend_bottom(ch)
    cat_axis(ch.x_axis, reverse=horizontal)
    hide_value_axis(ch.y_axis)
    style_chart_frame(ch)
    return ch


def composition_chart(cats_ref, series, labels=True, legend=True, ymax=None, title=None, axis_size=9):
    """Gestapelte Saeulen je Zeitpunkt. series: Liste (Werte, Titel-Zelle, Fuellfarbe, Textfarbe)."""
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.style = 2
    for i, (values_ref, title_ref, color, txt) in enumerate(series):
        s = ref_series(i, values_ref, title_ref, cats_ref)
        s.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill="FFFFFF", w=9525))
        s.invertIfNegative = False
        if labels:
            s.dLbls = data_labels(txt, pos="ctr", numfmt="0;-0;;", size=8)
        ch.series.append(s)
    ch.gapWidth = 45
    if legend:
        legend_bottom(ch, 8)
    else:
        ch.legend = None
    cat_axis(ch.x_axis, size=axis_size)
    if ymax:
        show_value_axis(ch.y_axis, ymax=ymax, major=10 if ymax > 20 else 5, size=7)
    else:
        hide_value_axis(ch.y_axis)
    if title:
        ch.title = chart_title(title, 9)
    style_chart_frame(ch)
    return ch


def verlauf_chart(cats_ref, values_ref, values_title_ref, soll_ref, soll_title_ref, title=None, color=C_PLAN):
    """Saeulen = Mitarbeiter auf der Stufe je Zeitpunkt, gestrichelte Linie = Soll (dieselbe Achse)."""
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.style = 2
    s = fill_series(ref_series(0, values_ref, values_title_ref, cats_ref), color)
    s.dLbls = data_labels(INK2, pos="outEnd", numfmt="0;-0;0", size=7)
    bar.series.append(s)
    bar.gapWidth = 50

    line = LineChart()
    ls = ref_series(1, soll_ref, soll_title_ref, cats_ref)
    ls.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=INK, w=19050, prstDash="dash"))
    ls.marker = Marker(symbol="none")
    ls.smooth = False
    line.series.append(ls)

    cat_axis(bar.x_axis, size=7)
    hide_value_axis(bar.y_axis)
    line.x_axis = bar.x_axis
    line.y_axis = bar.y_axis
    bar += line
    bar.legend = None
    if title:
        bar.title = chart_title(title, 9)
    style_chart_frame(bar)
    return bar


def place(ws, chart, top_left, bottom_right):
    """Diagramm genau auf einen Zellbereich legen (Ecke oben links von top_left bis Ecke oben links von bottom_right).
    So passt es unabhaengig von Schriftart/Spaltenbreite in Excel und LibreOffice."""
    c1, r1 = coordinate_from_string(top_left)
    c2, r2 = coordinate_from_string(bottom_right)
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=column_index_from_string(c1) - 1, row=r1 - 1),
                                 to=AnchorMarker(col=column_index_from_string(c2) - 1, row=r2 - 1))
    ws.add_chart(chart)


def fix_row_heights(ws, height=15):
    """Feste Zeilenhoehe fuer alle Zeilen ohne eigene Hoehe (gleiches Bild in Excel und LibreOffice)."""
    for r in range(1, ws.max_row + 2):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = height


# ---------------------------------------------------------------- Bausteine
GRID_COLS = 16   # Spalten B..Q
GRID_WIDTH = 10.5


def setup_grid(ws, hidden=()):
    ws.column_dimensions["A"].width = 2
    for i in range(2, 2 + GRID_COLS):
        ws.column_dimensions[get_column_letter(i)].width = GRID_WIDTH
    ws.column_dimensions["R"].width = 2
    for col in hidden:
        ws.column_dimensions[col].hidden = True
    for r in (1, 4):
        ws.row_dimensions[r].height = 8


def link_formula(target, text, cell="A1"):
    link = "#" + quote_sheetname(target).replace('"', '""') + "!" + cell
    return f'=HYPERLINK("{link}","{text}")'


def header_block(ws, title, subtitle_formula, links=()):
    """Titel (B2:H2), bis zu 4 Navigations-Links (I2:Q2), Untertitel (B3:Q3)."""
    merge_put(ws, "B2:H2", title, f=font(18, True), al=LEFT)
    ws.row_dimensions[2].height = 30
    slots = ["I2:J2", "K2:L2", "M2:N2", "O2:Q2"][-len(links):] if links else []
    for c, (target, text) in zip(slots, links):
        merge_put(ws, c, link_formula(target, text), f=font(10, color=LINK, underline="single"), al=RIGHT)
    merge_put(ws, "B3:Q3", subtitle_formula, f=font(9, color=INK2), al=LEFT)
    ws.row_dimensions[3].height = 16


def kpi_tiles(ws, tiles, row=5):
    """tiles: Liste (Titel, Wert-Formel, Unterzeile-Formel, Akzentfarbe); je 3 Spalten ab B."""
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 34
    ws.row_dimensions[row + 2].height = 16
    for i, (label, value, sub, accent) in enumerate(tiles):
        c1 = 2 + 3 * i
        c3 = c1 + 2
        L1, L3 = get_column_letter(c1), get_column_letter(c3)
        merge_put(ws, f"{L1}{row}:{L3}{row}", label, f=font(10, True, INK2),
                  al=Alignment(horizontal="left", vertical="bottom", indent=1))
        merge_put(ws, f"{L1}{row+1}:{L3}{row+1}", value, f=font(24, True, INK),
                  al=Alignment(horizontal="left", vertical="center", indent=1), nf="0")
        merge_put(ws, f"{L1}{row+2}:{L3}{row+2}", sub, f=font(9, color=MUTED),
                  al=Alignment(horizontal="left", vertical="top", indent=1))
        accent_side = Side(style="thick", color=accent)
        for r in range(row, row + 3):
            for c in range(c1, c3 + 1):
                cell = ws.cell(r, c)
                cell.fill = fill(TILE_BG)
                cell.border = Border(left=accent_side if c == c1 else None,
                                     right=Side(style="thick", color="FFFFFF") if c == c3 else None)


def explain(ws, row, text, height=40):
    """Kurze Erklaerung, wie gezaehlt wird (grauer Kasten ueber die ganze Breite)."""
    merge_put(ws, f"B{row}:Q{row}", text, f=font(9, color=INK2), fl=fill("F7F7F5"),
              al=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    ws.row_dimensions[row].height = height


def section_title(ws, ref_range, text):
    merge_put(ws, ref_range, text, f=font(12, True), al=LEFT)
    r = int("".join(ch for ch in ref_range.split(":")[0] if ch.isdigit()))
    ws.row_dimensions[r].height = 20


def section_note(ws, ref_range, text, height=None):
    merge_put(ws, ref_range, text, f=font(9, color=MUTED), al=LEFT_WRAP)
    if height:
        r = int("".join(ch for ch in ref_range.split(":")[0] if ch.isdigit()))
        ws.row_dimensions[r].height = height


def spacer(ws, *rows, height=8):
    for r in rows:
        ws.row_dimensions[r].height = height


def status_header_cells(ws, row, first_col):
    for code in range(1, 10):
        head(ws, f"{get_column_letter(first_col + code - 1)}{row}", f"={stat_label_ref(code)}")


def pct_sub(count_ref, total_ref, word="Mitarbeitern"):
    return f'=IF({total_ref}=0,"–",TEXT({count_ref}/{total_ref},"0%")&" von "&{total_ref}&" {word}")'


def compare_formula(ist, soll):
    return (f'=IF({ist}={soll},"Soll erreicht",IF({ist}<{soll},IF({soll}-{ist}=1,"es fehlt 1","es fehlen "&({soll}-{ist})),'
            f'({ist}-{soll})&" mehr als Soll"))')


def add_compare_cf(ws, cell_range, ist_first, soll_first):
    """Farben fuer Soll-Ist: rot = es fehlen, gelb = mehr als Soll, gruen = Soll erreicht.
    ist_first/soll_first: Zellen (relativ) fuer die linke obere Zelle von cell_range."""
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"{ist_first}<{soll_first}"], fill=fill(BAD_BG), font=Font(color=BAD_INK)))
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"{ist_first}>{soll_first}"], fill=fill(MORE_BG), font=Font(color=MORE_INK)))
    ws.conditional_formatting.add(cell_range, FormulaRule(
        formula=[f"{ist_first}={soll_first}"], fill=fill(GOOD_BG), font=Font(color=GOOD_INK)))


# ---------------------------------------------------------------- Blatt: Daten (Auswertung je Mitarbeiter)
FLAG_TITLES = ["Hinweis: Jahr vorbei ohne x", "Hinweis: weder x noch Jahr", "Hinweis: Ziel ohne Planung",
               "Hinweis: Ziel unklar", "Hinweis: heute ohne Stufe", "Hinweis: Ist-Qualifikation passt nicht",
               "Hinweis: Stufe fehlt", "Hinweis: doppelt"]


def _data_columns():
    cols = [("bez", "Bezirk (Blatt)", 15), ("row", "Zeile im Blatt", 7), ("name", "Name", 16), ("first", "Vorname", 14),
            ("ziel", "Ziel-Qualifikation (G)", 16), ("ist", "Ist-Qualifikation (J)", 16), ("ma", "Mitarbeiter (1 = ja)", 9)]
    cols += [(f"raw_{c}", f"Eintrag Spalte {c}", 8) for c in SRC_ENTRY]
    cols += [(f"v_{c}", f"Wert {c} (0 = x, Jahr, 9999 = leer)", 10) for c in SRC_ENTRY]
    cols += [(f"st{k}", None, 9) for k in range(N_TL)]
    cols += [(f"tl{k}", None, 9) for k in range(N_TL)]
    cols += [("zc", "Ziel-Spalte (1 = L … 8 = S, 0 = keine)", 10), ("vpc", "Verwendungsprüfung (1 = U, 2 = V)", 10),
             ("zgrp", "Ziel-Gruppe", 14), ("zv", "Wert in der Ziel-Spalte", 10), ("vpx", "„x“ bei Verwendungsprüfung", 10),
             ("status", "Status-Nr", 7), ("stext", "Status", 18), ("year", "Abschluss (Jahr in der Ziel-Spalte)", 10),
             ("ausb", "In Ausbildung (1 = ja)", 9), ("sort", "Sortierung", 9), ("rank", "Rang im Bezirk", 8),
             ("key", "Schlüssel", 18), ("stname", "Stufe heute (Text)", 18), ("istnr", "Ist-Qualifikation: Spalte", 9)]
    cols += [(f"f{i}", t, 28) for i, t in enumerate(FLAG_TITLES, start=1)]
    cols += [("hints", "Hinweise (alle)", 60), ("hflag", "Hinweis (1 = ja)", 8), ("hnr", "Nr. in der Prüfliste", 8)]
    return cols


DATA_COLS = _data_columns()
DC = {key: get_column_letter(i) for i, (key, _, _) in enumerate(DATA_COLS, start=1)}
DATA_LAST = 2              # letzte Zeile im Blatt Daten (wird in build() gesetzt)


def D(key):
    """Begrenzter Bereich einer Spalte im Blatt Daten (keine ganzen Spalten)."""
    c = DC[key]
    return f"Daten!${c}$2:${c}${DATA_LAST}"


def parse_formula(c):
    """Eintrag -> 0 ("x"), Jahr (27 -> 2027, 2027, Datum) oder 9999 (leer bzw. kein gueltiger Eintrag, z. B. 0)."""
    v = f"VALUE({c})"
    return (f'IF({c}="",9999,IF(LOWER({c})="x",0,IFERROR(IF({v}>2100,YEAR({v}),IF({v}>=1900,ROUND({v},0),'
            f'IF(AND({v}>=1,{v}<=99),2000+ROUND({v},0),9999))),9999)))')


def build_data(ws, districts):
    ws.sheet_view.showGridLines = True
    for key, title, width in DATA_COLS:
        col = DC[key]
        ws.column_dimensions[col].width = width
        if title is None:
            k = int(key[2:])
            what = "Stufe" if key.startswith("st") else "Teamleiter"
            unit = " (0 = keine, 1 = Azubi … 5 = RBEG)" if what == "Stufe" else " (1 = ja)"
            if k == 0:
                title = f"{what} heute{unit}"
            elif k == N_TL - 1:
                title = f"{what} nach Plan{unit}"
            else:
                title = f'="{what} Ende "&{year_ref(k)}&"{unit}"'
        put(ws, f"{col}1", title, f=font(9, True, "FFFFFF"), fl=fill("52514E"), al=LEFT_WRAP)
    ws.row_dimensions[1].height = 54
    ws.freeze_panes = "C2"
    last = DATA_LAST
    ws.auto_filter.ref = f"A1:{DC['hnr']}{last}"
    plain = font(9)
    c = DC
    r = 2
    for d in districts:
        s = quote_sheetname(d.sheet)

        def src(col):
            return f"{s}!${col}$1:${col}${SRC_MAXROW}"
        for k in range(SRC_ROWS):
            R = lambda key: f"{c[key]}{r}"
            f = {
                "bez": d.sheet,
                "row": SRC_FIRST_ROW + k,
                "name": f'=TRIM(CLEAN(INDEX({src(SRC_NAME)},${R("row")})&""))',
                "first": f'=TRIM(CLEAN(INDEX({src(SRC_FIRST)},${R("row")})&""))',
                "ziel": f'=TRIM(CLEAN(INDEX({src(SRC_ZIEL)},${R("row")})&""))',
                "ist": f'=TRIM(CLEAN(INDEX({src(SRC_IST)},${R("row")})&""))',
                "ma": f'=IF(AND({R("name")}<>"",{R("name")}<>"Name"),1,0)',
            }
            for col in SRC_ENTRY:
                f[f"raw_{col}"] = f'=TRIM(CLEAN(INDEX({src(col)},${R("row")})&""))'
                f[f"v_{col}"] = "=" + parse_formula(R(f"raw_{col}"))
            v = {col: R(f"v_{col}") for col in SRC_ENTRY}
            ma = f"${R('ma')}"
            # Stufe je Zeitpunkt: rechteste Spalte L..P mit "x" oder Jahr <= Zeitpunkt; "x" in U = Wmech, "x" in V = SigMech
            for k2 in range(N_TL):
                t = "0" if k2 == 0 else ("9998" if k2 == N_TL - 1 else year_ref(k2))
                parts = [f"({v[col]}<={t})*{i}" for i, (_, col) in enumerate(STAGES, start=1)]
                parts += [f"({v['U']}=0)*3", f"({v['V']}=0)*4"]
                f[f"st{k2}"] = f'=IF({ma}=0,"",MAX({",".join(parts)}))'
                f[f"tl{k2}"] = f'=IF({ma}=0,"",IF({v["S"]}<={t},1,0))'
            # Status an der Ziel-Qualifikation
            zc, vpc, zv, vpx, st = R("zc"), R("vpc"), R("zv"), R("vpx"), R("status")
            f["zc"] = f'=IF(OR({ma}=0,{R("ziel")}=""),0,IFERROR(INDEX({E_MAP_COLNO},MATCH({R("ziel")},{E_MAP_TEXT},0))*1,0))'
            f["vpc"] = f'=IF({zc}=0,0,IFERROR(INDEX({E_MAP_VPNO},MATCH({R("ziel")},{E_MAP_TEXT},0))*1,0))'
            f["zgrp"] = (f'=IF({zc}=0,"{OTHER}",IFERROR(IF(INDEX({E_MAP_CAT},MATCH({R("ziel")},{E_MAP_TEXT},0))&""="",'
                         f'"{OTHER}",INDEX({E_MAP_CAT},MATCH({R("ziel")},{E_MAP_TEXT},0))&""),"{OTHER}"))')
            f["zv"] = f'=IF({zc}=0,9999,INDEX({v["L"]}:{v["S"]},1,{zc}))'
            f["vpx"] = f'=IF({vpc}=1,IF({v["U"]}=0,1,0),IF({vpc}=2,IF({v["V"]}=0,1,0),0))'
            f["status"] = (f'=IF({ma}=0,"",IF({zc}=0,9,IF(OR({zv}=0,{vpx}=1),1,IF({zv}=9999,8,'
                           f'IF({zv}<{E_YEAR},7,MIN(6,2+{zv}-{E_YEAR}))))))')
            f["stext"] = f'=IF({ma}=0,"",INDEX({E_STAT_LABEL},{st}))'
            f["year"] = f'=IF({ma}=0,"",IF(AND({st}>=2,{st}<=7),{zv},""))'
            f["ausb"] = f'=IF({ma}=0,0,IF(AND({st}>=2,{st}<=6),1,0))'
            f["sort"] = f'=IF({ma}=0,"",{st}*1000+{R("row")})'
            f["rank"] = f'=IF({ma}=0,"",COUNTIFS({D("bez")},{R("bez")},{D("sort")},"<"&{R("sort")})+1)'
            f["key"] = f'=IF({ma}=0,"",{R("bez")}&"|"&{R("rank")})'
            st0, tl0 = R("st0"), R("tl0")
            f["stname"] = (f'=IF({ma}=0,"",IF({st0}=0,"–",INDEX({E_STAGE_NAMES},{st0}))&IF({tl0}=1," + Teamleiter",""))')
            f["istnr"] = (f'=IF(OR({ma}=0,{R("ist")}=""),0,IF(LOWER(LEFT({R("ist")},5))="azubi",1,'
                          f'IFERROR(INDEX({E_MAP_COLNO},MATCH({R("ist")},{E_MAP_TEXT},0))*1,0)))')
            # Hinweise fuer die Pruefliste
            past = "&".join(f'IF(AND({v[col]}>0,{v[col]}<{E_YEAR}),", {col} "&{v[col]},"")' for col in CHECK_COLS)
            f["f1"] = f'=IF({ma}=0,"",IF({past}="","","Jahr vorbei ohne „x“: "&MID({past},3,999)))'
            bad = "&".join(f'IF(AND({R("raw_" + col)}<>"",{v[col]}=9999),", {col} „"&{R("raw_" + col)}&"“","")'
                           for col in CHECK_COLS)
            f["f2"] = f'=IF({ma}=0,"",IF({bad}="","","Weder „x“ noch Jahr: "&MID({bad},3,999)))'
            f["f3"] = (f'=IF({st}=8,"Ziel „"&{R("ziel")}&"“: nichts geplant in Spalte "&CHAR(75+{zc})'
                       f'&IF({vpc}=1,"/U",IF({vpc}=2,"/V","")),"")')
            f["f4"] = (f'=IF({st}=9,IF({R("ziel")}="","Keine Ziel-Qualifikation eingetragen",'
                       f'"Ziel „"&{R("ziel")}&"“ hat keine eigene Spalte"),"")')
            f["f5"] = f'=IF({ma}=0,"",IF({st0}=0,"Heute auf keiner Stufe (kein „x“ in L–P)",""))'
            ist, istnr = R("ist"), R("istnr")
            f["f6"] = (f'=IF({ma}=0,"",IF(AND({istnr}>=1,{istnr}<=5,{istnr}<>{st0}),"Ist-Qualifikation „"&{ist}&"“, laut L–P heute "'
                       f'&IF({st0}=0,"keine Stufe",INDEX({E_STAGE_NAMES},{st0})),IF(AND({istnr}=8,{v["S"]}<>0),'
                       f'"Ist-Qualifikation „"&{ist}&"“, aber kein „x“ in S","")))')
            raw = lambda col: R("raw_" + col)
            gap = (f'IF(AND({raw("U")}<>"",{raw("N")}=""),"; U eingetragen, aber N leer","")'
                   f'&IF(AND({raw("V")}<>"",{raw("O")}=""),"; V eingetragen, aber O leer","")'
                   f'&IF(AND({raw("P")}<>"",{raw("O")}="",{v["V"]}<>0),"; P eingetragen, aber O leer","")')
            f["f7"] = f'=IF({ma}=0,"",IF({gap}="","","Stufe fehlt: "&MID({gap},3,999)))'
            dup = f'COUNTIFS({D("name")},{R("name")},{D("first")},{R("first")},{D("ma")},1)'
            f["f8"] = f'=IF({ma}=0,"",IF({dup}>1,"Name und Vorname stehen "&{dup}&"× in der Liste",""))'
            f["hints"] = "=MID(" + "&".join(f'IF({R(f"f{i}")}<>""," · "&{R(f"f{i}")},"")' for i in range(1, 9)) + ",4,9999)"
            f["hflag"] = f'=IF({R("hints")}<>"",1,0)'
            f["hnr"] = f'=IF({R("hflag")}=1,SUM(${c["hflag"]}$2:{R("hflag")}),"")'
            for key, val in f.items():
                cell = ws[f"{c[key]}{r}"]
                cell.value = val
                cell.font = plain
            r += 1
    return last


# ---------------------------------------------------------------- Bezirksseite: Zeilen
DR_EXPL = 9
DR_SEC1, DR_CH1, DR_CH1_END = 11, 12, 31
DR_SEC2, DR_NOTE2, DR_CH2, DR_CH2_END = 32, 33, 34, 49
DR_SEC3, DR_NOTE3, DR_CH3, DR_CH3_END = 50, 51, 52, 62
DR_T1 = 63                                  # Zahlen: Mitarbeiter je Stufe
DR_T1_HDR = DR_T1 + 1
DT_ROWS = {name: DR_T1_HDR + 1 + i for i, name in enumerate(STAGE_NAMES + [TL_NAME, NO_STAGE, "gesamt"])}
DR_T1_NOTE = DT_ROWS["gesamt"] + 1
DR_T2 = DR_T1_NOTE + 2                      # Abgleich mit dem Soll
DR_T2_HDR = DR_T2 + 1
DT_CMP = {cat: DR_T2_HDR + 1 + i for i, cat in enumerate(CATS)}
DT_MISSING = DT_CMP[CATS[-1]] + 1
DT_MORE = DT_MISSING + 1
DR_T3 = DT_MORE + 2                         # Status nach Ziel-Qualifikation
DR_T3_HDR = DR_T3 + 1
DT_ST = {g: DR_T3_HDR + 1 + i for i, g in enumerate(GROUPS)}
DT_ST_SUM = DT_ST[OTHER] + 1
DT_ST_PCT = DT_ST_SUM + 1
DR_LEG = DT_ST_PCT + 2                      # Was bedeuten die Status?
DR_LIST = DR_LEG + 1 + len(STATUS) + 2      # Mitarbeiterliste
DR_LIST_HDR = DR_LIST + 1
DR_LIST_FIRST = DR_LIST_HDR + 1
DR_LIST_LAST = DR_LIST_FIRST + LIST_ROWS - 1

T_SOLL = "D"
TBL = [get_column_letter(5 + k) for k in range(N_TL)]          # E..M: heute, 7 Jahre, nach Plan
T_PLAN = TBL[-1]
STAT_COL1 = 4                                                  # Spalte D = Status 1 ... L = Status 9
HELP0 = 22                                                     # Spalte V: Hilfswerte fuer Diagramme
HELP_SOLL = [get_column_letter(HELP0 + k) for k in range(N_TL - 1)]                     # V..AC: Soll-Linie
COMP = [get_column_letter(HELP0 + N_TL - 1 + k) for k in range(N_TL)]                    # AD..AL: Verlauf + Ziel


def soll_formula(cat):
    return f'=IF($T$2=0,0,IFERROR(INDEX({ZZ_AREA},MATCH("{cat}",{ZZ_LABELS},0),$T$2)*1,0))'


def stage_table(ws, hdr, soll_fn, count_fn, note):
    """Tabelle Mitarbeiter je Stufe: Spalten Soll, heute, Jahre, nach Plan. count_fn(row_name, k) -> Formel.
    Hilfswerte fuer Diagramme in ausgeblendeten Spalten (Soll-Linie, Verlauf + Ziel)."""
    head(ws, f"B{hdr}:C{hdr}", "Stufe", align="left")
    head(ws, f"{T_SOLL}{hdr}", "Soll (Zielzustand)")
    for k, col in enumerate(TBL):
        head(ws, f"{col}{hdr}", timeline_label(k))
    ws.row_dimensions[hdr].height = 30
    for k, col in enumerate(COMP[:-1]):
        put(ws, f"{col}{hdr}", f"={TBL[k]}{hdr}", f=font(8, color=MUTED))
    put(ws, f"{COMP[-1]}{hdr}", "Ziel", f=font(8, color=MUTED))
    rows = {name: hdr + 1 + i for i, name in enumerate(STAGE_NAMES + [TL_NAME, NO_STAGE, "gesamt"])}
    for name, r in rows.items():
        total = name == "gesamt"
        muted = name in ("Azubi", NO_STAGE)
        label = "Mitarbeiter gesamt" if total else ("Teamleiter (zusätzlich)" if name == TL_NAME else name)
        border = TOP_INK if total else BOTTOM_HAIR
        merge_put(ws, f"B{r}:C{r}", label, f=font(10, total or not muted, INK2 if muted else INK), al=LEFT, border=border)
        if name in CATS:
            soll = soll_fn(name)
        elif total:
            soll = "=" + "+".join(f"{T_SOLL}{rows[c]}" for c in CATS)
        else:
            soll = "–"
        put(ws, f"{T_SOLL}{r}", soll, f=font(10, True, MUTED if soll == "–" else INK), al=CENTER, nf="0", border=border)
        for k, col in enumerate(TBL):
            put(ws, f"{col}{r}", count_fn(name, k), f=font(10, total or col == T_PLAN, INK2 if muted else INK),
                al=CENTER, nf="0", border=border)
        if name in CATS:
            for col in HELP_SOLL:
                put(ws, f"{col}{r}", f"=${T_SOLL}{r}", f=font(8, color=MUTED), nf="0")
        if name in STACK_ORDER:
            for k, col in enumerate(COMP[:-1]):
                put(ws, f"{col}{r}", f"={TBL[k]}{r}", f=font(8, color=MUTED), nf="0")
            put(ws, f"{COMP[-1]}{r}", f"={T_SOLL}{r}" if name in CATS else 0, f=font(8, color=MUTED), nf="0")
    nr = rows["gesamt"] + 1
    merge_put(ws, f"B{nr}:Q{nr}", note, f=font(8, italic=True, color=MUTED), al=LEFT_WRAP)
    ws.row_dimensions[nr].height = 24
    return rows


def compare_table(ws, hdr, count_rows):
    """Abgleich je Stufe und Zeitpunkt: "Soll erreicht" / "es fehlen X" / "X mehr als Soll"."""
    head(ws, f"B{hdr}:C{hdr}", "Stufe", align="left")
    head(ws, f"{T_SOLL}{hdr}", "Soll (Zielzustand)")
    for k, col in enumerate(TBL):
        head(ws, f"{col}{hdr}", timeline_label(k))
    ws.row_dimensions[hdr].height = 30
    rows = {}
    for i, cat in enumerate(CATS):
        r = hdr + 1 + i
        rows[cat] = r
        cr = count_rows[cat]
        merge_put(ws, f"B{r}:C{r}", cat, f=font(10, True), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"{T_SOLL}{r}", f"={T_SOLL}{cr}", f=font(10, True), al=CENTER, nf="0", border=BOTTOM_HAIR)
        for col in TBL:
            put(ws, f"{col}{r}", compare_formula(f"{col}{cr}", f"${T_SOLL}{cr}"), f=font(9, col == T_PLAN),
                al=CENTER_WRAP, border=BOTTOM_HAIR)
        ws.row_dimensions[r].height = 26
    first = hdr + 1
    cr0 = count_rows[CATS[0]]
    add_compare_cf(ws, f"{TBL[0]}{first}:{T_PLAN}{first + len(CATS) - 1}", f"{TBL[0]}{cr0}", f"${T_SOLL}{cr0}")
    rm, rp = first + len(CATS), first + len(CATS) + 1
    merge_put(ws, f"B{rm}:C{rm}", "Fehlende Stellen", f=font(10, True, BAD_INK), al=LEFT, border=TOP_INK)
    merge_put(ws, f"B{rp}:C{rp}", "Mehr als Soll", f=font(10, True, MORE_INK), al=LEFT, border=BOTTOM_HAIR)
    put(ws, f"{T_SOLL}{rm}", "Summe", f=font(8, color=MUTED), al=CENTER, border=TOP_INK)
    put(ws, f"{T_SOLL}{rp}", "Summe", f=font(8, color=MUTED), al=CENTER, border=BOTTOM_HAIR)
    for col in TBL:
        miss = "+".join(f"MAX(0,${T_SOLL}{count_rows[c]}-{col}{count_rows[c]})" for c in CATS)
        more = "+".join(f"MAX(0,{col}{count_rows[c]}-${T_SOLL}{count_rows[c]})" for c in CATS)
        put(ws, f"{col}{rm}", "=" + miss, f=font(10, True, BAD_INK), al=CENTER, nf="0", border=TOP_INK)
        put(ws, f"{col}{rp}", "=" + more, f=font(10, True, MORE_INK), al=CENTER, nf="0", border=BOTTOM_HAIR)
    return rows, rm, rp


def comp_series(title, rows, keys=STACK_ORDER):
    out = []
    for key in keys:
        r = rows[key]
        color, txt = STAGE_COLORS[key]
        out.append((rng(title, COMP[0], r, COMP[-1], r), rng(title, "B", r), color, txt))
    return out


def status_legend(ws, row):
    section_title(ws, f"B{row}:Q{row}", "Was bedeuten die Status?")
    for code, color, _ in STATUS:
        r = row + code
        put(ws, f"B{r}", None, fl=fill(color))
        merge_put(ws, f"C{r}:E{r}", f"={stat_label_ref(code)}", f=font(10, True),
                  al=Alignment(horizontal="left", vertical="center", indent=1))
        merge_put(ws, f"F{r}:Q{r}", f"=Einstellungen!$E${E_STAT_FIRST + code - 1}", f=font(9, color=INK2), al=LEFT_WRAP)
        ws.row_dimensions[r].height = 18


# ---------------------------------------------------------------- Blatt: Dashboard je Bezirk
def build_district(ws, d):
    setup_grid(ws, hidden=["S", "T", "U"] + HELP_SOLL + COMP)
    title = ws.title
    put(ws, "S1", d.sheet)
    put(ws, "T2", f'=IFERROR(INDEX({E_BEZ_ZZCOL},MATCH($S$1,{E_BEZ_SHEET},0))*1,0)')
    header_block(ws, f"Ausbildungsstand – {d.sheet}",
                 f'="Bezugsjahr "&{E_YEAR}&"   ·   Quelle: Blatt „{d.sheet}“   ·   Soll: Zielzustand „"&IFERROR(INDEX({E_BEZ_ZZNAME},'
                 f'MATCH($S$1,{E_BEZ_SHEET},0)),"–")&"“   ·   Werte passen sich automatisch an, wenn das Blatt „{d.sheet}“ '
                 f'oder der Zielzustand geändert wird"',
                 links=[("Übersicht", "← Übersicht"), ("Soll-Ist", "Soll-Ist alle Bezirke"), (d.sheet, f"Blatt „{d.sheet}“")])

    st_sum = lambda code: f"{get_column_letter(STAT_COL1 + code - 1)}{DT_ST_SUM}"
    total = f"$M${DT_ST_SUM}"
    davon = (f'COUNTIFS({D("bez")},$S$1,{D("ausb")},1,{D("year")},{E_DAVON})')
    kpi_tiles(ws, [
        ("Mitarbeiter", f"={total}", f'=IF({st_sum(9)}>0,"davon "&{st_sum(9)}&" mit Ziel unklar","im Bezirk")', INK2),
        ("Fertig", f"={st_sum(1)}", pct_sub(st_sum(1), total), STATUS[0][1]),
        ("In Ausbildung", f"=SUM({st_sum(2)}:{st_sum(6)})",
         f'="davon "&{davon}&" mit Abschluss "&{E_DAVON}&IF({st_sum(7)}>0," · "&{st_sum(7)}&" überfällig","")', "2A78D6"),
        ("Fehlt (nichts geplant)", f"={st_sum(8)}", pct_sub(st_sum(8), total), STATUS[7][1]),
        ("Fehlende Stellen zum Soll", f"={TBL[0]}{DT_MISSING}",
         f'="Stand heute · nach Plan: "&{T_PLAN}{DT_MISSING}&" · Soll: "&{T_SOLL}{DT_ROWS["gesamt"]}&" Stellen"', BAD_INK),
    ])
    explain(ws, DR_EXPL, TXT_STATUS)
    spacer(ws, 8, 10, DR_CH1_END)

    # Zahlentabellen (zuerst, die Diagramme verweisen darauf)
    section_title(ws, f"B{DR_T1}:Q{DR_T1}", "Zahlen: Mitarbeiter je Stufe (zum Jahresende, laut Planung)")
    cnt = {name: i for i, name in enumerate(STAGE_NAMES, start=1)}

    def count(name, k):
        if name == "gesamt":
            return f'=COUNTIFS({D("bez")},$S$1,{D("ma")},1)'
        if name == TL_NAME:
            return f'=COUNTIFS({D("bez")},$S$1,{D(f"tl{k}")},1)'
        return f'=COUNTIFS({D("bez")},$S$1,{D(f"st{k}")},{cnt.get(name, 0)})'
    rows = stage_table(ws, DR_T1_HDR, soll_formula, count,
                       "Teamleiter zählen zusätzlich bei ihrer Stufe (z. B. SigMech RBEG) – deshalb ist „Mitarbeiter gesamt“ "
                       "kleiner als die Summe der Zeilen. „ohne Stufe“ = noch kein „x“ (bzw. kein Jahr bis dahin) in L–P.")
    assert rows == DT_ROWS
    section_title(ws, f"B{DR_T2}:Q{DR_T2}", "Abgleich mit dem Soll je Stufe")
    cmp_rows, rm, rp = compare_table(ws, DR_T2_HDR, rows)
    assert cmp_rows == DT_CMP and (rm, rp) == (DT_MISSING, DT_MORE)
    spacer(ws, DR_T1_NOTE + 1, DT_MORE + 1)

    section_title(ws, f"B{DR_T3}:Q{DR_T3}", "Zahlen: Status nach Ziel-Qualifikation")
    head(ws, f"B{DR_T3_HDR}:C{DR_T3_HDR}", "Ziel-Qualifikation", align="left")
    status_header_cells(ws, DR_T3_HDR, STAT_COL1)
    head(ws, f"M{DR_T3_HDR}", "Summe")
    ws.row_dimensions[DR_T3_HDR].height = 30
    for g, r in DT_ST.items():
        merge_put(ws, f"B{r}:C{r}", g if g != OTHER else "Sonstige / ohne Ziel", al=LEFT, border=BOTTOM_HAIR)
        for code in range(1, 10):
            col = get_column_letter(STAT_COL1 + code - 1)
            put(ws, f"{col}{r}", f'=COUNTIFS({D("bez")},$S$1,{D("zgrp")},"{g}",{D("status")},{code})',
                al=CENTER, nf=NF_COUNT, border=BOTTOM_HAIR)
        put(ws, f"M{r}", f"=SUM(D{r}:L{r})", f=font(10, True), al=CENTER, nf="0", border=BOTTOM_HAIR)
    merge_put(ws, f"B{DT_ST_SUM}:C{DT_ST_SUM}", "Summe", f=font(10, True), al=LEFT)
    merge_put(ws, f"B{DT_ST_PCT}:C{DT_ST_PCT}", "Anteil", f=font(9, color=INK2), al=LEFT)
    for ci in range(STAT_COL1, 14):
        col = get_column_letter(ci)
        first, last = DT_ST[GROUPS[0]], DT_ST[OTHER]
        put(ws, f"{col}{DT_ST_SUM}", f"=SUM({col}{first}:{col}{last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)
        put(ws, f"{col}{DT_ST_PCT}", f'=IF({total}=0,"",{col}{DT_ST_SUM}/{total})', f=font(9, color=INK2), al=CENTER, nf="0%")
    spacer(ws, DT_ST_PCT + 1)
    status_legend(ws, DR_LEG)
    spacer(ws, DR_LEG + len(STATUS) + 1)

    # Diagramme Seite 1: Status | Soll / heute / nach Plan je Stufe
    section_title(ws, f"B{DR_SEC1}:I{DR_SEC1}", "Mitarbeiter nach Status")
    section_title(ws, f"J{DR_SEC1}:Q{DR_SEC1}", "Soll, heute und nach Plan je Stufe")
    place(ws, status_bar_chart(ws, DT_ST_SUM, STAT_COL1, STAT_COL1 + 8, DR_T3_HDR), f"B{DR_CH1}", f"J{DR_CH1_END}")
    c1, c2 = rows[CATS[0]], rows[CATS[-1]]
    cats = rng(title, "B", c1, "B", c2)
    place(ws, grouped_chart(cats, [
        (rng(title, T_SOLL, c1, T_SOLL, c2), rng(title, T_SOLL, DR_T1_HDR), C_SOLL),
        (rng(title, TBL[0], c1, TBL[0], c2), rng(title, TBL[0], DR_T1_HDR), C_HEUTE),
        (rng(title, T_PLAN, c1, T_PLAN, c2), rng(title, T_PLAN, DR_T1_HDR), C_PLAN),
    ]), f"J{DR_CH1}", f"R{DR_CH1_END}")

    # Seite 2: Verlauf gestapelt + kleine Verlaufsdiagramme
    section_title(ws, f"B{DR_SEC2}:Q{DR_SEC2}", "Verlauf je Jahr: Mitarbeiter je Stufe – rechts daneben der Zielzustand")
    explain(ws, DR_NOTE2, TXT_COUNT, height=40)
    comp_cats = rng(title, COMP[0], DR_T1_HDR, COMP[-1], DR_T1_HDR)
    place(ws, composition_chart(comp_cats, comp_series(title, rows)), f"B{DR_CH2}", f"R{DR_CH2_END}")
    spacer(ws, DR_CH2_END)
    section_title(ws, f"B{DR_SEC3}:Q{DR_SEC3}", "Verlauf je Stufe gegen das Soll")
    section_note(ws, f"B{DR_NOTE3}:Q{DR_NOTE3}",
                 "Säulen = Mitarbeiter auf dieser Stufe (heute und zum Jahresende laut Planung) · gestrichelte Linie = Soll (Zielzustand)")
    tl_cats = rng(title, TBL[0], DR_T1_HDR, TBL[-2], DR_T1_HDR)
    anchors = ["B", "E", "H", "K", "N", "Q"]
    for k, cat in enumerate(CATS):
        r = rows[cat]
        ch = verlauf_chart(tl_cats, rng(title, TBL[0], r, TBL[-2], r), rng(title, "B", r),
                           rng(title, HELP_SOLL[0], r, HELP_SOLL[-1], r), rng(title, T_SOLL, DR_T1_HDR), title=cat)
        place(ws, ch, f"{anchors[k]}{DR_CH3}", f"{anchors[k + 1]}{DR_CH3_END}")
    spacer(ws, DR_CH3_END)

    # Mitarbeiterliste
    section_title(ws, f"B{DR_LIST}:Q{DR_LIST}", "Mitarbeiterliste (sortiert nach Status)")
    cols = [("B", "B", "Nr."), ("C", "D", "Name"), ("E", "F", "Vorname"), ("G", "H", "Ziel-Qualifikation"),
            ("I", "J", "Ist-Qualifikation"), ("K", "L", "Stufe heute (L–P)"), ("M", "M", "Abschluss"), ("N", "Q", "Status")]
    for a, b, text in cols:
        head(ws, f"{a}{DR_LIST_HDR}" if a == b else f"{a}{DR_LIST_HDR}:{b}{DR_LIST_HDR}", text, align="left")
    ws.row_dimensions[DR_LIST_HDR].height = 18
    for k in range(1, LIST_ROWS + 1):
        r = DR_LIST_FIRST + k - 1
        idx = f"$S{r}"
        get = lambda key: f'=IF({idx}="","",INDEX({D(key)},{idx}))'
        put(ws, f"S{r}", f'=IFERROR(MATCH($S$1&"|"&{k},{D("key")},0),"")')
        put(ws, f"T{r}", f'=IF({idx}="","",INDEX({D("status")},{idx}))')
        put(ws, f"B{r}", f'=IF({idx}="","",{k})', f=font(9, color=MUTED), al=LEFT)
        merge_put(ws, f"C{r}:D{r}", get("name"), f=font(10, True), al=LEFT)
        merge_put(ws, f"E{r}:F{r}", get("first"), al=LEFT)
        merge_put(ws, f"G{r}:H{r}", get("ziel"), al=LEFT)
        merge_put(ws, f"I{r}:J{r}", get("ist"), f=font(10, color=INK2), al=LEFT)
        merge_put(ws, f"K{r}:L{r}", get("stname"), al=LEFT)
        put(ws, f"M{r}", f'=IF({idx}="","",IF(INDEX({D("year")},{idx})="","–",INDEX({D("year")},{idx})))', al=CENTER, nf="0")
        merge_put(ws, f"N{r}:Q{r}", get("stext"), f=font(10, True), al=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[r].height = 14
    over = DR_LIST_LAST + 1
    merge_put(ws, f"B{over}:Q{over}",
              f'=IF({total}>{LIST_ROWS},"Hinweis: Der Bezirk hat mehr als {LIST_ROWS} Mitarbeiter – die Liste zeigt nur die ersten {LIST_ROWS}.","")',
              f=font(9, True, STATUS[7][1]), al=LEFT)
    # Bereiche ueberlappen nicht: LibreOffice wendet je Zelle nur einen Bereich an
    ws.conditional_formatting.add(f"B{DR_LIST_FIRST}:M{DR_LIST_LAST}",
                                  FormulaRule(formula=[f'$S{DR_LIST_FIRST}<>""'], border=BOTTOM_HAIR))
    chip_border = Border(bottom=Side(style="thin", color="FFFFFF"))
    for code, color, txt in STATUS:
        ws.conditional_formatting.add(f"N{DR_LIST_FIRST}:Q{DR_LIST_LAST}", FormulaRule(
            formula=[f"$T{DR_LIST_FIRST}={code}"], fill=fill(color), font=Font(color=txt, bold=True), border=chip_border))
    # Druckseiten: Ueberblick | Verlauf | Zahlen Soll-Ist | Status | Mitarbeiterliste
    page_setup(ws, f"A1:Q{over}", breaks=[DR_CH1_END, DR_CH3_END, DT_MORE + 1, DR_LIST - 1])
    fix_row_heights(ws)
    ws.sheet_properties.tabColor = "86B6EF"
    return rows


# ---------------------------------------------------------------- Blatt: Uebersicht (Status)
OV_EXPL = 9
OV_SEC1, OV_CH1, OV_CH1_END = 11, 12, 31
OV_T = 32
OV_HDR = OV_T + 1
OV_FIRST = OV_HDR + 1


def build_overview(ws, districts):
    setup_grid(ws, hidden=("S",))
    n = len(districts)
    ov_last = OV_FIRST + n - 1
    ov_sum, ov_pct = ov_last + 1, ov_last + 2
    header_block(ws, "Ausbildungsstand LST – alle Bezirke",
                 f'="Bezugsjahr "&{E_YEAR}&"   ·   {n} Bezirke   ·   Klick auf einen Bezirk in der Tabelle öffnet seine Seite   ·   '
                 f'Werte passen sich automatisch an, wenn die Bezirks-Blätter geändert werden"',
                 links=[("Soll-Ist", "Soll-Ist-Vergleich →"), ("Prüfliste", "Prüfliste →"), ("Einstellungen", "Einstellungen →")])
    sc = lambda code: f"{get_column_letter(4 + code)}{ov_sum}"   # Status 1 -> Spalte E
    total = f"$N${ov_sum}"
    davon = f'COUNTIFS({D("ausb")},1,{D("year")},{E_DAVON})'
    kpi_tiles(ws, [
        ("Mitarbeiter gesamt", f"={total}", f'=IF({sc(9)}>0,"davon "&{sc(9)}&" mit Ziel unklar","in allen Bezirken")', INK2),
        ("Fertig", f"={sc(1)}", pct_sub(sc(1), total), STATUS[0][1]),
        ("In Ausbildung", f"=SUM({sc(2)}:{sc(6)})",
         f'="davon "&{davon}&" mit Abschluss "&{E_DAVON}&IF({sc(7)}>0," · "&{sc(7)}&" überfällig","")', "2A78D6"),
        ("Fehlt (nichts geplant)", f"={sc(8)}", pct_sub(sc(8), total), STATUS[7][1]),
        ("Fehlende Stellen zum Soll", f"='Soll-Ist'!{SI_TB_SUM_REF[0]}", f'="Stand heute · nach Plan: "&\'Soll-Ist\'!{SI_TB_SUM_REF[1]}',
         BAD_INK),
    ])
    explain(ws, OV_EXPL, TXT_STATUS)
    spacer(ws, 8, 10, OV_CH1_END)
    section_title(ws, f"B{OV_SEC1}:G{OV_SEC1}", "Alle Bezirke nach Status")
    section_title(ws, f"H{OV_SEC1}:Q{OV_SEC1}", "Status je Bezirk")

    section_title(ws, f"B{OV_T}:Q{OV_T}", "Zahlen je Bezirk")
    head(ws, f"B{OV_HDR}:D{OV_HDR}", "Bezirk (Klick öffnet die Seite)", align="left")
    status_header_cells(ws, OV_HDR, 5)
    head(ws, f"N{OV_HDR}", "Summe")
    head(ws, f"O{OV_HDR}:P{OV_HDR}", "Anteil fertig")
    head(ws, f"Q{OV_HDR}", "Hinweise (Prüfliste)")
    ws.row_dimensions[OV_HDR].height = 30
    for i, d in enumerate(districts):
        r = OV_FIRST + i
        put(ws, f"S{r}", d.sheet)
        merge_put(ws, f"B{r}:D{r}", link_formula(dash_name(d.sheet), d.sheet), f=font(10, True, LINK, underline="single"),
                  al=LEFT, border=BOTTOM_HAIR)
        for code in range(1, 10):
            col = get_column_letter(4 + code)
            put(ws, f"{col}{r}", f'=COUNTIFS({D("bez")},$S{r},{D("status")},{code})', al=CENTER, nf=NF_COUNT, border=BOTTOM_HAIR)
        put(ws, f"N{r}", f"=SUM(E{r}:M{r})", f=font(10, True), al=CENTER, nf="0", border=BOTTOM_HAIR)
        merge_put(ws, f"O{r}:P{r}", f'=IF(N{r}=0,"",E{r}/N{r})', al=CENTER, nf="0%", border=BOTTOM_HAIR)
        put(ws, f"Q{r}", f'=COUNTIFS({D("bez")},$S{r},{D("hflag")},1)', al=CENTER, nf=NF_COUNT, border=BOTTOM_HAIR)
    merge_put(ws, f"B{ov_sum}:D{ov_sum}", "Summe", f=font(10, True), al=LEFT)
    merge_put(ws, f"B{ov_pct}:D{ov_pct}", "Anteil", f=font(9, color=INK2), al=LEFT)
    for ci in range(5, 15):
        col = get_column_letter(ci)
        put(ws, f"{col}{ov_sum}", f"=SUM({col}{OV_FIRST}:{col}{ov_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)
        put(ws, f"{col}{ov_pct}", f'=IF({total}=0,"",{col}{ov_sum}/{total})', f=font(9, color=INK2), al=CENTER, nf="0%")
    merge_put(ws, f"O{ov_sum}:P{ov_sum}", f'=IF({total}=0,"",E{ov_sum}/{total})', f=font(10, True), al=CENTER, nf="0%",
              border=TOP_INK)
    put(ws, f"Q{ov_sum}", f"=SUM(Q{OV_FIRST}:Q{ov_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)

    place(ws, status_bar_chart(ws, ov_sum, 5, 13, OV_HDR), f"B{OV_CH1}", f"H{OV_CH1_END}")
    place(ws, stacked_status_chart(ws, OV_HDR, OV_FIRST, ov_last, 2, 5, 13), f"H{OV_CH1}", f"R{OV_CH1_END}")

    lg = ov_pct + 2
    spacer(ws, ov_pct + 1)
    status_legend(ws, lg)
    foot = lg + len(STATUS) + 2
    merge_put(ws, f"B{foot}:Q{foot}",
              "Bezugsjahr und Zuordnungen: Blatt „Einstellungen“ · Soll ändern: Blatt „Zielzustand“ · auffällige Einträge: "
              "Blatt „Prüfliste“ · Auswertung je Mitarbeiter: Blatt „Daten“ (ganz hinten).",
              f=font(9, italic=True, color=MUTED), al=LEFT)
    page_setup(ws, f"A1:Q{foot}", breaks=[OV_CH1_END])
    fix_row_heights(ws)
    ws.sheet_properties.tabColor = "104281"


# ---------------------------------------------------------------- Blatt: Soll-Ist (alle Bezirke)
SI_SEC1, SI_CH1, SI_CH1_END = 11, 12, 31
SI_SEC2, SI_NOTE2, SI_SM = 32, 33, 34
SM_COLS, SM_H = 4, 9                   # kleine Diagramme: 4 je Zeile, 9 Zeilen hoch
SI_SM_BANDS = 3
SI_TG = SI_SM + SI_SM_BANDS * (SM_H + 1)          # Zahlen LST gesamt
SI_TG_HDR = SI_TG + 1
SI_TG_ROWS = {name: SI_TG_HDR + 1 + i for i, name in enumerate(STAGE_NAMES + [TL_NAME, NO_STAGE, "gesamt"])}
SI_TG_MISS = SI_TG_ROWS["gesamt"] + 2             # (Zeile davor: Hinweis aus stage_table)
SI_TG_MORE = SI_TG_MISS + 1
SI_TB = SI_TG_MORE + 2                            # fehlende Stellen je Bezirk
SI_TB_HDR = SI_TB + 1
SI_TB_FIRST = SI_TB_HDR + 1
SI_TB_SUM_REF = ("$E$0", "$M$0")                  # wird in build() gesetzt


def build_soll_ist(ws, districts, ymax):
    setup_grid(ws, hidden=["S", "T", "U"] + HELP_SOLL + COMP)
    title = ws.title
    n = len(districts)
    gsheets = [dash_name(d.sheet) for d in districts]
    tb_last = SI_TB_FIRST + n - 1
    tb_sum = tb_last + 1
    tq = tb_sum + 3
    tq_h1, tq_h2 = tq + 2, tq + 3
    tq_first, tq_last = tq_h2 + 1, tq_h2 + n
    header_block(ws, "Soll-Ist-Vergleich LST – alle Bezirke",
                 f'="Bezugsjahr "&{E_YEAR}&"   ·   Soll = Spalte „Zielzustand“ (Blatt Zielzustand)   ·   jeder Mitarbeiter zählt einmal – '
                 f'auf seiner Stufe in L–P   ·   „Nach Plan“ = alle geplanten Ausbildungen abgeschlossen"',
                 links=[("Übersicht", "← Übersicht"), ("Zielzustand", "Soll ändern →")])

    def g_sum(row, col):
        return "=" + "+".join(f"{quote_sheetname(g)}!${col}${row}" for g in gsheets)

    # Zahlen LST gesamt (Summe der Bezirke)
    section_title(ws, f"B{SI_TG}:Q{SI_TG}", "Zahlen: LST gesamt – Mitarbeiter je Stufe (Summe aller Bezirke)")
    rows = stage_table(ws, SI_TG_HDR, lambda cat: g_sum(DT_ROWS[cat], T_SOLL), lambda name, k: g_sum(DT_ROWS[name], TBL[k]),
                       "Summe der Bezirksseiten. Teamleiter zählen zusätzlich bei ihrer Stufe. Fehlende Stellen und „mehr als Soll“ "
                       "werden je Bezirk und Stufe gerechnet und dann addiert – ein Überschuss in einem Bezirk gleicht keinen "
                       "fehlenden Platz in einem anderen aus.")
    assert rows == SI_TG_ROWS
    for r, label, src, color in ((SI_TG_MISS, "Fehlende Stellen", DT_MISSING, BAD_INK), (SI_TG_MORE, "Mehr als Soll", DT_MORE, MORE_INK)):
        merge_put(ws, f"B{r}:C{r}", label, f=font(10, True, color), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"{T_SOLL}{r}", "Summe", f=font(8, color=MUTED), al=CENTER, border=BOTTOM_HAIR)
        for col in TBL:
            put(ws, f"{col}{r}", g_sum(src, col), f=font(10, True, color), al=CENTER, nf="0", border=BOTTOM_HAIR)
    spacer(ws, SI_TG_MORE + 1)

    # Fehlende Stellen je Bezirk
    section_title(ws, f"B{SI_TB}:Q{SI_TB}", "Zahlen: fehlende Stellen je Bezirk (alle Stufen zusammen)")
    head(ws, f"B{SI_TB_HDR}:C{SI_TB_HDR}", "Bezirk (Klick öffnet die Seite)", align="left")
    head(ws, f"{T_SOLL}{SI_TB_HDR}", "Soll (Stellen)")
    for k, col in enumerate(TBL):
        head(ws, f"{col}{SI_TB_HDR}", timeline_label(k))
    head(ws, f"N{SI_TB_HDR}:O{SI_TB_HDR}", "Mehr als Soll (nach Plan)")
    head(ws, f"P{SI_TB_HDR}:Q{SI_TB_HDR}", "Mitarbeiter")
    ws.row_dimensions[SI_TB_HDR].height = 30
    for i, d in enumerate(districts):
        r = SI_TB_FIRST + i
        g = quote_sheetname(gsheets[i])
        merge_put(ws, f"B{r}:C{r}", link_formula(gsheets[i], d.sheet), f=font(10, True, LINK, underline="single"),
                  al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"{T_SOLL}{r}", f"={g}!${T_SOLL}${DT_ROWS['gesamt']}", f=font(10, True), al=CENTER, nf="0", border=BOTTOM_HAIR)
        for col in TBL:
            put(ws, f"{col}{r}", f"={g}!${col}${DT_MISSING}", f=font(10, col in (TBL[0], T_PLAN)), al=CENTER, nf=NF_COUNT,
                border=BOTTOM_HAIR)
        merge_put(ws, f"N{r}:O{r}", f"={g}!${T_PLAN}${DT_MORE}", al=CENTER, nf=NF_COUNT, border=BOTTOM_HAIR)
        merge_put(ws, f"P{r}:Q{r}", f"={g}!${TBL[0]}${DT_ROWS['gesamt']}", al=CENTER, nf="0", border=BOTTOM_HAIR)
    merge_put(ws, f"B{tb_sum}:C{tb_sum}", "Summe", f=font(10, True), al=LEFT, border=TOP_INK)
    for col in [T_SOLL] + TBL:
        put(ws, f"{col}{tb_sum}", f"=SUM({col}{SI_TB_FIRST}:{col}{tb_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)
    merge_put(ws, f"N{tb_sum}:O{tb_sum}", f"=SUM(N{SI_TB_FIRST}:N{tb_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)
    merge_put(ws, f"P{tb_sum}:Q{tb_sum}", f"=SUM(P{SI_TB_FIRST}:P{tb_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)
    ws.conditional_formatting.add(f"{TBL[0]}{SI_TB_FIRST}:{T_PLAN}{tb_last}", FormulaRule(
        formula=[f"{TBL[0]}{SI_TB_FIRST}>0"], fill=fill(BAD_BG), font=Font(color=BAD_INK)))
    ws.conditional_formatting.add(f"{TBL[0]}{SI_TB_FIRST}:{T_PLAN}{tb_last}", FormulaRule(
        formula=[f"{TBL[0]}{SI_TB_FIRST}=0"], fill=fill(GOOD_BG), font=Font(color=GOOD_INK)))
    spacer(ws, tb_sum + 1, tb_sum + 2)

    # Soll / heute / nach Plan je Bezirk und Stufe
    section_title(ws, f"B{tq}:Q{tq}", "Zahlen: Soll, heute und nach Plan je Bezirk und Stufe")
    section_note(ws, f"B{tq + 1}:Q{tq + 1}", "grün = Soll erreicht · rot = es fehlen Mitarbeiter · gelb = mehr als Soll")
    head(ws, f"B{tq_h1}:B{tq_h2}", "Bezirk", align="left")
    for j, cat in enumerate(CATS):
        cs, ch_, cp = (get_column_letter(3 + 3 * j + i) for i in range(3))
        head(ws, f"{cs}{tq_h1}:{cp}{tq_h1}", cat)
        for col, text in zip((cs, ch_, cp), ("Soll", "heute", "Plan")):
            head(ws, f"{col}{tq_h2}", text)
    for i, d in enumerate(districts):
        r = tq_first + i
        g = quote_sheetname(gsheets[i])
        put(ws, f"B{r}", d.sheet, f=font(9, True), al=Alignment(horizontal="left", vertical="center", shrink_to_fit=True),
            border=BOTTOM_HAIR)
        for j, cat in enumerate(CATS):
            cs, ch_, cp = (get_column_letter(3 + 3 * j + i2) for i2 in range(3))
            put(ws, f"{cs}{r}", f"={g}!${T_SOLL}${DT_ROWS[cat]}", f=font(10, True), al=CENTER, nf="0", border=BOTTOM_HAIR)
            put(ws, f"{ch_}{r}", f"={g}!${TBL[0]}${DT_ROWS[cat]}", al=CENTER, nf="0", border=BOTTOM_HAIR)
            put(ws, f"{cp}{r}", f"={g}!${T_PLAN}${DT_ROWS[cat]}", al=CENTER, nf="0", border=BOTTOM_HAIR)
    for j in range(len(CATS)):
        cs, ch_, cp = (get_column_letter(3 + 3 * j + i2) for i2 in range(3))
        add_compare_cf(ws, f"{ch_}{tq_first}:{cp}{tq_last}", f"{ch_}{tq_first}", f"${cs}{tq_first}")
    put(ws, f"B{tq_last + 1}", "Summe", f=font(10, True), al=LEFT, border=TOP_INK)
    for ci in range(3, 3 + 3 * len(CATS)):
        col = get_column_letter(ci)
        put(ws, f"{col}{tq_last + 1}", f"=SUM({col}{tq_first}:{col}{tq_last})", f=font(10, True), al=CENTER, nf="0", border=TOP_INK)

    # Kennzahlen
    gr = SI_TG_ROWS
    kpi_tiles(ws, [
        ("Soll (Zielzustand)", f"={T_SOLL}{gr['gesamt']}", f'="Stellen in "&{n}&" Bezirken"', INK2),
        ("Mitarbeiter heute", f"={TBL[0]}{gr['gesamt']}", f'="davon "&{TBL[0]}{gr["Azubi"]}&" Azubis, "&{TBL[0]}{gr[NO_STAGE]}&" ohne Stufe"',
         C_HEUTE),
        ("Fehlende Stellen heute", f"={TBL[0]}{tb_sum}", '="Summe aller Bezirke und Stufen"', BAD_INK),
        ("Fehlende Stellen nach Plan", f"={T_PLAN}{tb_sum}", '="wenn alle Planungen umgesetzt sind"', C_PLAN),
        ("Mehr als Soll nach Plan", f"=N{tb_sum}", f'="davon Arbeiter LST: "&{T_PLAN}{gr["Arbeiter LST"]}', "EDA100"),
    ])
    explain(ws, 9, TXT_COUNT)
    spacer(ws, 8, 10, SI_CH1_END)

    # Diagramme
    section_title(ws, f"B{SI_SEC1}:H{SI_SEC1}", "Fehlende Stellen je Bezirk")
    section_title(ws, f"I{SI_SEC1}:Q{SI_SEC1}", "LST gesamt: Mitarbeiter je Stufe – rechts der Zielzustand")
    cats = rng(title, "B", SI_TB_FIRST, "B", tb_last)
    place(ws, grouped_chart(cats, [
        (rng(title, TBL[0], SI_TB_FIRST, TBL[0], tb_last), rng(title, TBL[0], SI_TB_HDR), C_HEUTE),
        (rng(title, T_PLAN, SI_TB_FIRST, T_PLAN, tb_last), rng(title, T_PLAN, SI_TB_HDR), C_PLAN),
    ], horizontal=True, label_size=8), f"B{SI_CH1}", f"I{SI_CH1_END}")
    comp_cats = rng(title, COMP[0], SI_TG_HDR, COMP[-1], SI_TG_HDR)
    place(ws, composition_chart(comp_cats, comp_series(title, rows)), f"I{SI_CH1}", f"R{SI_CH1_END}")

    section_title(ws, f"B{SI_SEC2}:Q{SI_SEC2}", "Verlauf je Bezirk gegen das Soll")
    section_note(ws, f"B{SI_NOTE2}:Q{SI_NOTE2}",
                 "Je Bezirk: Mitarbeiter je Stufe heute und zum Jahresende, rechts der Zielzustand · Farben wie im Diagramm oben · "
                 "alle Diagramme mit gleicher Skala")
    for i, d in enumerate(districts):
        band, pos = divmod(i, SM_COLS)
        top = SI_SM + band * (SM_H + 1)
        c0 = 2 + 4 * pos
        g = gsheets[i]
        dc = rng(g, COMP[0], DR_T1_HDR, COMP[-1], DR_T1_HDR)
        ch = composition_chart(dc, comp_series(g, DT_ROWS), labels=False, legend=False, ymax=ymax, title=d.sheet, axis_size=7)
        place(ws, ch, f"{get_column_letter(c0)}{top}", f"{get_column_letter(c0 + 4)}{top + SM_H}")
        spacer(ws, top + SM_H)
    # Druckseiten: Kennzahlen + Diagramme | Verlauf je Bezirk | Zahlen | Tabelle je Stufe
    page_setup(ws, f"A1:Q{tq_last + 1}", breaks=[SI_CH1_END, SI_TG - 1, tq - 1])
    fix_row_heights(ws)
    ws.sheet_properties.tabColor = "0CA30C"
    return tb_sum


# ---------------------------------------------------------------- Blatt: Pruefliste
PR_HDR = 7
PR_FIRST = PR_HDR + 1


def build_pruefliste(ws):
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 5, "C": 13, "D": 15, "E": 13, "F": 14, "G": 14, "H": 92, "I": 2}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["J"].hidden = True
    merge_put(ws, "B2:G2", "Prüfliste: auffällige Einträge", f=font(18, True), al=LEFT)
    put(ws, "H2", link_formula("Übersicht", "← Übersicht"), f=font(10, color=LINK, underline="single"), al=RIGHT)
    ws.row_dimensions[2].height = 30
    merge_put(ws, "B3:H3", f'=SUM({D("hflag")})&" Mitarbeiter mit Hinweisen · Bezugsjahr "&{E_YEAR}&'
                           f'" · die Liste aktualisiert sich automatisch, wenn die Bezirks-Blätter geändert werden"',
              f=font(9, color=INK2), al=LEFT)
    merge_put(ws, "B4:H4",
              "Geprüft wird je Mitarbeiter: Jahr vorbei ohne „x“ (Spalten L–P, S, U, V) · Eintrag, der weder „x“ noch ein Jahr ist "
              "(z. B. 0) · Ziel ohne passende Planung (Status „Fehlt“) · Ziel unklar · heute auf keiner Stufe (kein „x“ in L–P) · "
              "Ist-Qualifikation (J) passt nicht zu den Kreuzen in L–P/S · Verwendungsprüfung oder RBEG eingetragen, aber die Stufe "
              "davor ist leer · Name und Vorname mehrfach. Korrigiert wird im jeweiligen Bezirks-Blatt.",
              f=font(9, color=INK2), fl=fill("F7F7F5"), al=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    ws.row_dimensions[4].height = 42
    ws.row_dimensions[5].height = 8
    for col, text in zip("BCDEFGH", ["Nr.", "Bezirk", "Name", "Vorname", "Ziel-Qualifikation", "Ist-Qualifikation", "Hinweise"]):
        head(ws, f"{col}{PR_HDR}", text, align="left")
    ws.row_dimensions[PR_HDR].height = 18
    for k in range(1, PRUEF_ROWS + 1):
        r = PR_FIRST + k - 1
        idx = f"$J{r}"
        get = lambda key: f'=IF({idx}="","",INDEX({D(key)},{idx}))'
        put(ws, f"J{r}", f'=IFERROR(MATCH({k},{D("hnr")},0),"")')
        put(ws, f"B{r}", f'=IF({idx}="","",{k})', f=font(9, color=MUTED), al=Alignment(vertical="top"))
        for col, key, bold in (("C", "bez", False), ("D", "name", True), ("E", "first", False), ("F", "ziel", False), ("G", "ist", False)):
            put(ws, f"{col}{r}", get(key), f=font(10, bold), al=TOP_WRAP)
        put(ws, f"H{r}", get("hints"), f=font(10), al=TOP_WRAP)
        ws.row_dimensions[r].height = 26
    over = PR_FIRST + PRUEF_ROWS
    merge_put(ws, f"B{over}:H{over}",
              f'=IF(SUM({D("hflag")})>{PRUEF_ROWS},"Hinweis: mehr als {PRUEF_ROWS} Mitarbeiter mit Hinweisen – die Liste zeigt die ersten {PRUEF_ROWS}.","")',
              f=font(9, True, STATUS[7][1]), al=LEFT)
    ws.conditional_formatting.add(f"B{PR_FIRST}:H{over - 1}", FormulaRule(formula=[f'$J{PR_FIRST}<>""'], border=BOTTOM_HAIR))
    ws.freeze_panes = f"A{PR_FIRST}"
    page_setup(ws, f"A1:H{over}", title_rows=f"{PR_HDR}:{PR_HDR}")
    ws.sheet_properties.tabColor = "EC835A"


# ---------------------------------------------------------------- Blatt: Zielzustand (Kopie)
def build_zielzustand(ws, src_ws, src_name):
    """Kopiert Tabelle1 aus LST_Zielzustand_gesamt.xlsx (Werte, Formeln, Formate)."""
    for row in src_ws.iter_rows():
        for c in row:
            if c.__class__.__name__ == "MergedCell":
                continue
            d = ws.cell(c.row, c.column)
            d.value = c.value
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.number_format = c.number_format
    for m in src_ws.merged_cells.ranges:
        ws.merge_cells(str(m))
    for k, dim in src_ws.column_dimensions.items():
        ws.column_dimensions[k].width = dim.width
        ws.column_dimensions[k].hidden = dim.hidden
    for k, dim in src_ws.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[k].height = dim.height
    ws.freeze_panes = "B3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 3             # breite Tabelle: 3 Seiten nebeneinander, Spalte A wiederholt
    ws.page_setup.fitToHeight = 1
    ws.print_title_cols = "A:A"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    r = src_ws.max_row + 2
    notes = [
        f"Quelle: {src_name}, Blatt „{src_ws.title}“ (übernommen). Hier das Soll ändern – alle Dashboard-Seiten passen sich an.",
        "Verwendet wird je Bezirk die Spalte „Zielzustand“ in den Zeilen Arbeiter LST, Wmech, SigMech, SigMech RBEG und Teamleiter. "
        "FBÜB und Azubi-Zeilen werden nicht verglichen.",
        "Welcher Bezirk zu welchem Blatt gehört, steht im Blatt „Einstellungen“.",
    ]
    for i, n in enumerate(notes):
        put(ws, f"A{r + i}", n, f=font(10, italic=True, color=INK2))
    put(ws, f"A{r + len(notes) + 1}", link_formula("Übersicht", "← Übersicht"), f=font(10, color=LINK, underline="single"))
    ws.sheet_properties.tabColor = "FAB219"


# ---------------------------------------------------------------- Blatt: Einstellungen
def build_settings(ws, districts):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [2, 28, 24, 15, 26, 18, 16, 60]):
        ws.column_dimensions[col].width = w
    put(ws, "B1", "Einstellungen & Erklärung", f=font(16, True))
    put(ws, "B2", "Gelb hinterlegte Zellen dürfen geändert werden. Alle Zahlen und Diagramme passen sich automatisch an.",
        f=font(10, color=INK2))
    put(ws, "H1", link_formula("Übersicht", "← Übersicht"), f=font(10, color=LINK, underline="single"), al=RIGHT)

    put(ws, "B4", "Bezugsjahr (aktuelles Jahr)", f=font(10, True), al=LEFT)
    put(ws, "C4", "=YEAR(TODAY())", f=font(11, True, "0000FF"), fl=fill(INPUT_BG), al=CENTER, nf="0")
    merge_put(ws, "D4:H4", "Standard: =JAHR(HEUTE()). Kann mit einer festen Zahl überschrieben werden (z. B. 2027). Geplante Jahre "
                           "davor gelten als „Überfällig“; der Verlauf zeigt 7 Jahre ab dem Bezugsjahr.", f=font(9, color=INK2), al=LEFT_WRAP)
    put(ws, "B5", "Jahr für „In Ausbildung, davon …“", f=font(10, True), al=LEFT)
    put(ws, "C5", "=C4+1", f=font(11, True, "0000FF"), fl=fill(INPUT_BG), al=CENTER, nf="0")
    merge_put(ws, "D5:H5", "Standard: Bezugsjahr + 1. Die Kennzahl „In Ausbildung“ zeigt, wie viele davon in diesem Jahr ihren "
                           "Abschluss haben.", f=font(9, color=INK2), al=LEFT_WRAP)
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 28

    put(ws, f"B{E_STAT_FIRST - 2}", "Status je Mitarbeiter (gemessen an der Ziel-Qualifikation, Spalte G)", f=font(12, True))
    for col, text in zip("BCDE", ["Nr.", "Bezeichnung", "Farbe", "Bedeutung"]):
        put(ws, f"{col}{E_STAT_FIRST - 1}", text, f=font(10, True, INK2), fl=fill(HEAD_BG), al=LEFT)
    for col in "FGH":
        put(ws, f"{col}{E_STAT_FIRST - 1}", None, fl=fill(HEAD_BG))
    for code, color, _ in STATUS:
        r = E_STAT_FIRST + code - 1
        label, meaning = STATUS_TEXT[code]
        if label is None:
            off = code - 2
            label = f'="Abschluss "&({E_YEAR}+{off})' if code < 6 else f'="Abschluss ab "&({E_YEAR}+{off})'
        put(ws, f"B{r}", code, al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"C{r}", label, f=font(10, True), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"D{r}", None, fl=fill(color), border=Border(bottom=Side(style="thin", color="FFFFFF")))
        merge_put(ws, f"E{r}:H{r}", meaning, f=font(9, color=INK2), al=LEFT_WRAP, border=BOTTOM_HAIR)
        ws.row_dimensions[r].height = 26 if len(meaning) > 100 else 16

    put(ws, f"B{E_STAGE_FIRST - 2}", "Stufen für Soll-Ist und Verlauf (nur Spalten L–P, von links nach rechts)", f=font(12, True))
    for col, text in zip("BCDE", ["Nr.", "Stufe", "Spalte", "zählt auch bei „x“ in"]):
        put(ws, f"{col}{E_STAGE_FIRST - 1}", text, f=font(10, True, INK2), fl=fill(HEAD_BG), al=LEFT)
    for i, (name, col) in enumerate(STAGES, start=1):
        r = E_STAGE_FIRST + i - 1
        put(ws, f"B{r}", i, al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"C{r}", name, f=font(10, True), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"D{r}", col, al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"E{r}", {"N": "U (örtl. Verwendungsprüfung Wmech)", "O": "V (örtl. Verwendungsprüfung SigMech)"}.get(col, "–"),
            f=font(9, color=INK2), al=LEFT, border=BOTTOM_HAIR)
    put(ws, f"B{E_STAGE_LAST + 1}", "Teamleiter werden zusätzlich über Spalte S gezählt.", f=font(9, color=INK2))

    put(ws, f"B{E_BEZ_HDR - 2}", "Bezirke: Blatt → Bezirk im Zielzustand", f=font(12, True))
    put(ws, f"B{E_BEZ_HDR - 1}", "Name genau wie in Zeile 1 des Blatts „Zielzustand“. Das Soll steht dort in der Spalte „Zielzustand“ "
                                 "(3 Spalten rechts vom Namen).", f=font(9, color=INK2))
    for col, text in zip("BCDEF", ["Bezirks-Blatt", "Name im Zielzustand", "Spalte Soll (Nr., automatisch)",
                                   "Prüfung Zielzustand", "Prüfung Spalten im Blatt"]):
        put(ws, f"{col}{E_BEZ_HDR}", text, f=font(10, True, INK2), fl=fill(HEAD_BG), al=LEFT_WRAP)
    ws.row_dimensions[E_BEZ_HDR].height = 30
    for i in range(E_BEZ_LAST - E_BEZ_FIRST + 1):
        r = E_BEZ_FIRST + i
        d = districts[i] if i < len(districts) else None
        put(ws, f"B{r}", d.sheet if d else None, f=font(10, True), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"C{r}", d.zz_name if d else None, f=font(10, color="0000FF"), fl=fill(INPUT_BG), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"D{r}", f'=IF(C{r}="","",IFERROR(MATCH(C{r},Zielzustand!$A$1:$BZ$1,0)+3,""))', al=CENTER, border=BOTTOM_HAIR)
        put(ws, f"E{r}", f'=IF(B{r}="","",IF(D{r}="","Name nicht gefunden – kein Soll",IF(INDEX(Zielzustand!$A$2:$BZ$2,1,D{r})'
                         f'="Zielzustand","ok","Spalte „Zielzustand“ prüfen")))', al=LEFT, border=BOTTOM_HAIR)
        if d:
            s = quote_sheetname(d.sheet)
            checks = [f'COUNTIF({s}!${c}$1:${c}$2,"{t}")>0' for c, t in HEADERS.items()]
            checks += [f'OR(LEFT({s}!${c}$1,4)="{t}",LEFT({s}!${c}$2,4)="{t}")' for c, t in HEADER_PREFIX.items()]
            put(ws, f"F{r}", f'=IF(AND({",".join(checks)}),"ok","Spalten prüfen!")', al=LEFT, border=BOTTOM_HAIR)
        else:
            put(ws, f"F{r}", None, border=BOTTOM_HAIR)
    ws.conditional_formatting.add(f"E{E_BEZ_FIRST}:F{E_BEZ_LAST}", FormulaRule(
        formula=[f'AND(E{E_BEZ_FIRST}<>"",E{E_BEZ_FIRST}<>"ok")'], fill=fill(BAD_BG), font=Font(color=BAD_INK, bold=True)))

    put(ws, f"B{E_MAP_HDR - 2}", "Zuordnung Ziel-/Ist-Qualifikation → Spalte (für den Status)", f=font(12, True))
    put(ws, f"B{E_MAP_HDR - 1}", "Wird in Spalte G oder J ein neuer Begriff verwendet, hier eine Zeile ergänzen (Groß-/Kleinschreibung egal). "
                                 "Ohne Spalte → Status „Ziel unklar“.", f=font(9, color=INK2))
    for col, text in zip("BCDEFGH", ["Text in Spalte G / J", "Spalte (L–S)", "Nr. (auto)", "Örtl. Verwendungs­prüfung (U/V)",
                                     "Nr. (auto)", "Gruppe (Tabelle Status)", "Hinweis"]):
        put(ws, f"{col}{E_MAP_HDR}", text, f=font(10, True, INK2), fl=fill(HEAD_BG), al=LEFT_WRAP)
    ws.row_dimensions[E_MAP_HDR].height = 30
    hints = {"Azubi": "Ziel Azubi = Spalte L. Bei der Ist-Qualifikation zählt jeder Text, der mit „Azubi“ beginnt.",
             "Senior Expert LST": "keine eigene Spalte → Status „Ziel unklar“",
             "Weichmech": "fertig bei „x“ in N oder „x“ bei der örtl. Verwendungsprüfung Wmech (U)",
             "Sigmech": "fertig bei „x“ in O oder „x“ bei der örtl. Verwendungsprüfung SigMech (V)"}
    for i in range(E_MAP_LAST - E_MAP_FIRST + 1):
        r = E_MAP_FIRST + i
        text, col, vp, cat = MAPPING[i] if i < len(MAPPING) else (None, None, None, None)
        put(ws, f"B{r}", text, f=font(10, color="0000FF"), fl=fill(INPUT_BG), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"C{r}", col or None, f=font(10, color="0000FF"), fl=fill(INPUT_BG), al=CENTER, border=BOTTOM_HAIR)
        put(ws, f"D{r}", f'=IF(C{r}="","",IFERROR(IF(AND(CODE(UPPER(C{r}))>=76,CODE(UPPER(C{r}))<=83),CODE(UPPER(C{r}))-75,""),""))',
            al=CENTER, border=BOTTOM_HAIR)
        put(ws, f"E{r}", vp or None, f=font(10, color="0000FF"), fl=fill(INPUT_BG), al=CENTER, border=BOTTOM_HAIR)
        put(ws, f"F{r}", f'=IF(E{r}="","",IFERROR(IF(AND(CODE(UPPER(E{r}))>=85,CODE(UPPER(E{r}))<=86),CODE(UPPER(E{r}))-84,""),""))',
            al=CENTER, border=BOTTOM_HAIR)
        put(ws, f"G{r}", cat, f=font(10, color="0000FF"), fl=fill(INPUT_BG), al=LEFT, border=BOTTOM_HAIR)
        put(ws, f"H{r}", hints.get(text), f=font(9, color=INK2), al=LEFT, border=BOTTOM_HAIR)

    r = E_MAP_LAST + 3
    put(ws, f"B{r}", "So wird gezählt", f=font(12, True))
    notes = [
        f"Ausgewertet werden die Bezirks-Blätter aus der Tabelle „Bezirke“. Je Blatt werden die Zeilen {SRC_FIRST_ROW}–"
        f"{SRC_FIRST_ROW + SRC_ROWS - 1} gelesen. Eine Zeile zählt als Mitarbeiter, wenn in Spalte B ein Name steht (Kopfzeile „Name“ zählt nicht).",
        "Einträge in L–S, U, V: „x“ = fertig, Zahl = geplantes Jahr (27 = 2027). Alles andere (z. B. 0 oder ?) zählt wie ein leeres Feld "
        "und steht in der Prüfliste.",
        TXT_COUNT,
        "Fehlende Stellen = je Bezirk und Stufe: Soll minus Ist, wenn das Ist kleiner ist; dann addiert. Mehr als Soll = Ist minus Soll, "
        "wenn das Ist größer ist.",
        TXT_STATUS,
        "Das Blatt „Daten“ (ganz hinten) enthält die Auswertung je Mitarbeiter. Es wird per Formel erzeugt und sollte nicht von Hand "
        "geändert werden.",
    ]
    for i, n in enumerate(notes):
        merge_put(ws, f"B{r + 1 + i}:H{r + 1 + i}", "• " + n, f=font(10, color=INK2), al=LEFT_WRAP)
        ws.row_dimensions[r + 1 + i].height = 30 if len(n) < 250 else 58
    page_setup(ws, f"A1:H{r + len(notes)}", breaks=[E_STAGE_LAST + 1, E_BEZ_LAST + 1, r - 1])
    ws.sheet_properties.tabColor = "FAB219"


# ---------------------------------------------------------------- Hauptfunktion
def build(districts, zz_ws, zz_name, out_path, ymax=30):
    global DATA_LAST, SI_TB_SUM_REF
    DATA_LAST = 1 + len(districts) * SRC_ROWS
    tb_sum = SI_TB_FIRST + len(districts)
    SI_TB_SUM_REF = (f"${TBL[0]}${tb_sum}", f"${T_PLAN}${tb_sum}")
    wb = Workbook()
    ws_over = wb.active
    ws_over.title = "Übersicht"
    ws_si = wb.create_sheet("Soll-Ist")
    for d in districts:
        build_district(wb.create_sheet(dash_name(d.sheet)), d)
    assert build_soll_ist(ws_si, districts, ymax) == tb_sum
    build_overview(ws_over, districts)
    build_pruefliste(wb.create_sheet("Prüfliste"))
    build_zielzustand(wb.create_sheet("Zielzustand"), zz_ws, zz_name)
    build_settings(wb.create_sheet("Einstellungen"), districts)
    ws_data = wb.create_sheet("Daten")
    build_data(ws_data, districts)
    ws_data.sheet_properties.tabColor = "C3C2B7"
    wb.save(out_path)
    return [ws.title for ws in wb.worksheets]
