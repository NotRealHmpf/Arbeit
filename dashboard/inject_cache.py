"""Schreibt die von LibreOffice berechneten Ergebnisse als gespeicherte Werte in die
Formelzellen und Diagramme der neuen Blaetter. So zeigt die Datei die Zahlen auch dort,
wo nicht neu gerechnet wird (geschuetzte Ansicht, Vorschau). Excel rechnet beim Oeffnen
ohnehin neu (fullCalcOnLoad).
"""
import re
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CNS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
N, C = "{%s}" % NS, "{%s}" % CNS
ERRORS = {"#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"}


def dump(el):
    return etree.tostring(el, xml_declaration=True, encoding="UTF-8", standalone=True)


def fmt_num(v):
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int) or (isinstance(v, float) and v.is_integer()):
        return str(int(v))
    return repr(float(v))


def set_cell_cache(c, value):
    v = c.find(N + "v")
    if v is None:
        v = etree.SubElement(c, N + "v")
    if value is None:
        value = ""
    if isinstance(value, bool):
        c.set("t", "b")
        v.text = "1" if value else "0"
    elif isinstance(value, (int, float)):
        c.attrib.pop("t", None)
        v.text = fmt_num(value)
    elif isinstance(value, str) and value in ERRORS:
        c.set("t", "e")
        v.text = value
    else:
        c.set("t", "str")
        v.text = str(value)


def ref_values(values_wb, ref):
    sheet, rng = ref.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    ws = values_wb[sheet]
    c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
    out = []
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            out.append(ws.cell(r, c).value)
    return out


def fill_chart_cache(chart_xml, values_wb):
    root = etree.fromstring(chart_xml)
    for ref_el in list(root.iter(C + "numRef", C + "strRef")):
        f = ref_el.find(C + "f")
        vals = ref_values(values_wb, f.text)
        for old in ref_el.findall(C + "numCache") + ref_el.findall(C + "strCache"):
            ref_el.remove(old)
        is_text = any(isinstance(v, str) for v in vals)
        if ref_el.tag == C + "numRef" and is_text:
            ref_el.tag = C + "strRef"           # Textkategorien als strRef
        if ref_el.tag == C + "strRef":
            cache = etree.SubElement(ref_el, C + "strCache")
            etree.SubElement(cache, C + "ptCount").set("val", str(len(vals)))
            for i, v in enumerate(vals):
                if v is None:
                    continue
                pt = etree.SubElement(cache, C + "pt")
                pt.set("idx", str(i))
                etree.SubElement(pt, C + "v").text = str(v)
        else:
            cache = etree.SubElement(ref_el, C + "numCache")
            etree.SubElement(cache, C + "formatCode").text = "General"
            etree.SubElement(cache, C + "ptCount").set("val", str(len(vals)))
            for i, v in enumerate(vals):
                if not isinstance(v, (int, float)):
                    continue
                pt = etree.SubElement(cache, C + "pt")
                pt.set("idx", str(i))
                etree.SubElement(pt, C + "v").text = fmt_num(v)
    return dump(root)


def inject(merged_path, calc_path, out_path, sheet_names):
    values_wb = load_workbook(calc_path, data_only=True)
    zin = zipfile.ZipFile(merged_path)
    parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    order = [i.filename for i in zin.infolist()]

    wb = etree.fromstring(parts["xl/workbook.xml"])
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    target = {r.get("Id"): r.get("Target") for r in rels}
    sheet_parts = {}
    for sh in wb.find(N + "sheets"):
        if sh.get("name") in sheet_names:
            t = target[sh.get("{%s}id" % RNS)]
            sheet_parts[sh.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t

    n_cells = 0
    charts = set()
    for name, part in sheet_parts.items():
        ws = values_wb[name]
        root = etree.fromstring(parts[part])
        for c in root.iter(N + "c"):
            if c.find(N + "f") is None:
                continue
            set_cell_cache(c, ws[c.get("r")].value)
            n_cells += 1
        parts[part] = dump(root)
        # Diagramme dieses Blatts finden (Blatt-Rels -> Drawing -> Drawing-Rels -> Charts)
        srels = part.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
        if srels in parts:
            for r in etree.fromstring(parts[srels]):
                t = r.get("Target")
                if "drawings/drawing" in t:
                    d = t.lstrip("/")
                    drels = d.replace("xl/drawings/", "xl/drawings/_rels/") + ".rels"
                    for dr in etree.fromstring(parts[drels]):
                        charts.add(dr.get("Target").lstrip("/"))
    for ch in charts:
        parts[ch] = fill_chart_cache(parts[ch], values_wb)

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in order:
            zout.writestr(name, parts[name])
    return {"cells": n_cells, "charts": len(charts)}
