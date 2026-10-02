"""Erzeugt aus der PowerPoint „Kontingentierung EL Produkte“ eine Excel-Liste fuer einen
Standort: je Kontingent-Platz eine Zeile, in die ein Mitarbeiter eingetragen wird.

    python build_kontingente.py <kontingentierung.pptx> <ausgabe.xlsx> [STANDORT]

STANDORT ist eine Spalte der Folientabellen (FMZ, FFM, KO, KSL, RIS, P3), Standard FFM.
Grundlagencoachings werden nicht beruecksichtigt. Termine ohne Kontingent fuer den
Standort (z. B. „Kontingentierung Zentrale“, „Termin folgt“) stehen als Hinweis im
Blatt Uebersicht.
Benoetigt: python-pptx, openpyxl, lxml, LibreOffice (soffice) mit Calc.
"""
import os
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "dashboard"))
from build_dashboard import libreoffice_recalc  # noqa: E402
from inject_cache import ERRORS, inject  # noqa: E402

IGNORE = ("grundlagencoaching",)
DATE_RANGE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})\s*[-–]\s*(\d{1,2})\.(\d{1,2})\.(\d{2,4})")

# ---------------------------------------------------------------- Gestaltung
FONT = "Arial"
INK = "1B1F24"
MUTED = "6B7280"
REPEAT = "9AA1AC"          # wiederholter Seminar-/Terminname ab dem 2. Platz
HEAD_BG = "1F3A5F"
TILE_BG = "F1F4F8"
BAND_BG = "F6F8FB"         # jeder zweite Termin
INPUT_BG = "FFF4D6"        # auszufuellen
LINE = "D5DAE1"
GROUP_LINE = "8A97A8"
OPEN_BG, OPEN_FG = "FDE7D9", "B4521C"
DONE_BG, DONE_FG = "DDF1E2", "1E6B35"
PART_BG, PART_FG = "FFF1C2", "8A6A00"
DUP_BG, DUP_FG = "F8D2D2", "A11D1D"


@dataclass
class Termin:
    seminar: str
    termine: list          # ["23.08.2027 – 03.09.2027", ...]
    plaetze: int


@dataclass
class Hinweis:
    seminar: str
    termin: str
    text: str


def clean(text):
    return " ".join(text.split())


def fmt_ranges(text):
    out = []
    for d1, m1, y1, d2, m2, y2 in DATE_RANGE.findall(text):
        y1, y2 = (y if len(y) == 4 else "20" + y for y in (y1, y2))
        out.append(f"{int(d1):02d}.{int(m1):02d}.{y1} – {int(d2):02d}.{int(m2):02d}.{y2}")
    return out


def read_pptx(path, standort):
    termine, hinweise, ignoriert = [], [], []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            rows = [[clean(c.text) for c in r.cells] for r in shape.table.rows]
            head = rows[0]
            if standort not in head:
                continue
            col = head.index(standort)
            note_col = next((i for i, h in enumerate(head) if h.lower().startswith("mitarbeiter")), None)
            for row in rows[1:]:
                name, datum = row[0], row[1]
                if not name:
                    continue
                if any(k in name.lower() for k in IGNORE):
                    if name not in ignoriert:
                        ignoriert.append(name)
                    continue
                ranges = fmt_ranges(datum)
                quota = row[col]
                if quota.isdigit() and int(quota) > 0 and ranges:
                    termine.append(Termin(name, ranges, int(quota)))
                    continue
                note = row[note_col] if note_col is not None else ""
                if not note:
                    note = datum if not ranges else f"kein Kontingent für {standort}"
                hinweise.append(Hinweis(name, " u. ".join(ranges), note))
    return termine, hinweise, ignoriert


# ---------------------------------------------------------------- Excel
def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", fgColor=color)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def build(termine, hinweise, ignoriert, standort, jahr, out):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    ws = wb.active
    ws.title = f"Kontingente {standort}"
    ov = wb.create_sheet("Übersicht")
    main = f"'{ws.title}'"
    for sh in (ws, ov):
        sh.sheet_view.showGridLines = False
        sh.sheet_view.zoomScale = 100

    # ------------------------------------------------ Blatt Kontingente
    widths = {"A": 2.5, "B": 6, "C": 30, "D": 25, "E": 25, "F": 9, "G": 32, "H": 14, "I": 36}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w

    ws["B1"] = f"Kontingentierung EL-Produkte {jahr} – Standort {standort}"
    ws["B1"].font = font(16, True, HEAD_BG)
    ws["B2"] = (f"Quelle: Kontingentierung EL Produkte Mitte (Basisseminare / Entstörtraining). "
                f"Nur Kontingente {standort}, Grundlagencoachings nicht berücksichtigt.")
    ws["B2"].font = font(9, color=MUTED)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[3].height = 8

    first = 9
    last = first + sum(t.plaetze for t in termine) - 1
    col = {k: f"${k}${first}:${k}${last}" for k in "BCDEFGHI"}

    # Kennzahlen
    tiles = [("B4:C4", "B5:C5", "Plätze gesamt", f"=COUNTA({col['C']})", "0"),
             ("D4", "D5", "Eingeplant", f'=COUNTIF({col["H"]},"eingeplant")', "0"),
             ("E4", "E5", "Offen", "=B5-D5", "0"),
             ("F4:G4", "F5:G5", "Belegung", "=IF(B5=0,0,D5/B5)", "0%")]
    for lab_rng, val_rng, label, formula, numfmt in tiles:
        for rng in (lab_rng, val_rng):
            if ":" in rng:
                ws.merge_cells(rng)
        lab, val = lab_rng.split(":")[0], val_rng.split(":")[0]
        ws[lab], ws[val] = label, formula
        ws[lab].font = font(9, color=MUTED)
        ws[val].font = font(18, True, HEAD_BG)
        ws[val].number_format = numfmt
        for rng in (lab_rng, val_rng):
            for row in ws[rng if ":" in rng else f"{rng}:{rng}"]:
                for c in row:
                    c.fill = fill(TILE_BG)
                    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 30

    ws["B6"] = ("So geht’s: In jeder Zeile einen Mitarbeiter in der gelben Spalte „Mitarbeiter“ eintragen "
                "(z. B. „Müller, Thomas“). Status, Kennzahlen und Übersicht aktualisieren sich automatisch. "
                "Doppelte Einträge im selben Termin werden rot markiert.")
    ws.merge_cells("B6:I6")
    ws["B6"].font = font(9, italic=True, color=MUTED)
    ws["B6"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[6].height = 26
    ws.row_dimensions[7].height = 6

    heads = ["Nr.", "Seminar", "Termin 1", "Termin 2", "Platz", "Mitarbeiter", "Status", "Bemerkung"]
    for i, h in enumerate(heads):
        c = ws.cell(8, 2 + i, h)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(HEAD_BG)
        c.alignment = Alignment(horizontal="center" if h in ("Nr.", "Platz", "Status") else "left",
                                vertical="center", indent=0 if h in ("Nr.", "Platz", "Status") else 1)
    ws.row_dimensions[8].height = 24

    r = first
    for ti, t in enumerate(termine):
        new_seminar = ti == 0 or termine[ti - 1].seminar != t.seminar
        band = fill(BAND_BG) if ti % 2 else None
        for p in range(1, t.plaetze + 1):
            vals = [r - first + 1, t.seminar, t.termine[0], t.termine[1] if len(t.termine) > 1 else None,
                    f"{p} / {t.plaetze}", None, f'=IF(TRIM(G{r})="","offen","eingeplant")', None]
            for i, v in enumerate(vals):
                c = ws.cell(r, 2 + i, v)
                c.font = font(10, color=INK if p == 1 or i not in (1, 2, 3) else REPEAT, bold=(p == 1 and i == 1))
                c.alignment = Alignment(vertical="center",
                                        horizontal="center" if i in (0, 4, 6) else "left",
                                        indent=0 if i in (0, 4, 6) else 1)
                if i == 5:
                    c.fill = fill(INPUT_BG)
                elif band:
                    c.fill = band
                top = side(GROUP_LINE, "medium") if (p == 1 and new_seminar and ti) else \
                    side(GROUP_LINE) if p == 1 and ti else side()
                c.border = Border(top=top, bottom=side(), left=side(LINE) if i == 5 else None,
                                  right=side(LINE) if i == 5 else None)
            ws.cell(r, 8).font = font(10, True)
            ws.row_dimensions[r].height = 20
            r += 1
    for i in range(8):
        c = ws.cell(last, 2 + i)
        c.border = Border(top=c.border.top, bottom=side(GROUP_LINE, "medium"),
                          left=c.border.left, right=c.border.right)

    data = f"B{first}:I{last}"
    ws.conditional_formatting.add(
        f"G{first}:G{last}",
        FormulaRule(formula=[f'AND(TRIM($G{first})<>"",COUNTIFS({col["C"]},$C{first},{col["D"]},$D{first},'
                             f'{col["G"]},$G{first})>1)'],
                    fill=fill(DUP_BG), font=Font(name=FONT, color=DUP_FG, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"H{first}:H{last}",
        FormulaRule(formula=[f'$H{first}="offen"'], fill=fill(OPEN_BG), font=Font(name=FONT, color=OPEN_FG, bold=True)))
    ws.conditional_formatting.add(
        f"H{first}:H{last}",
        FormulaRule(formula=[f'$H{first}="eingeplant"'], fill=fill(DONE_BG), font=Font(name=FONT, color=DONE_FG, bold=True)))

    ws.auto_filter.ref = f"B8:I{last}"
    ws.freeze_panes = f"A{first}"
    ws.print_title_rows = "8:8"
    ws.print_area = f"A1:I{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Seite &P von &N"

    # ------------------------------------------------ Blatt Uebersicht
    for c, w in {"A": 2.5, "B": 30, "C": 25, "D": 25, "E": 12, "F": 12, "G": 10, "H": 14}.items():
        ov.column_dimensions[c].width = w
    ov["B1"] = f"Übersicht je Termin – Standort {standort}"
    ov["B1"].font = font(16, True, HEAD_BG)
    ov["B2"] = f"Zählt die eingetragenen Mitarbeiter aus dem Blatt „{ws.title}“."
    ov["B2"].font = font(9, color=MUTED)
    ov.row_dimensions[1].height = 26
    ov.row_dimensions[3].height = 8

    heads = ["Seminar", "Termin 1", "Termin 2", f"Plätze {standort}", "Eingeplant", "Offen", "Status"]
    for i, h in enumerate(heads):
        c = ov.cell(4, 2 + i, h)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(HEAD_BG)
        c.alignment = Alignment(horizontal="left" if i < 3 else "center", vertical="center", indent=1 if i < 3 else 0)
    ov.row_dimensions[4].height = 24

    o_first = 5
    for ti, t in enumerate(termine):
        r = o_first + ti
        new_seminar = ti == 0 or termine[ti - 1].seminar != t.seminar
        vals = [t.seminar, t.termine[0], t.termine[1] if len(t.termine) > 1 else None,
                t.plaetze,
                f'=COUNTIFS({main}!{col["C"]},$B{r},{main}!{col["D"]},$C{r},{main}!{col["H"]},"eingeplant")',
                f"=E{r}-F{r}",
                f'=IF(G{r}=0,"komplett",IF(F{r}=0,"offen","teilweise"))']
        for i, v in enumerate(vals):
            c = ov.cell(r, 2 + i, v)
            c.font = font(10, bold=(i == 0 and new_seminar) or i == 6,
                          color=INK if (i != 0 or new_seminar) else REPEAT)
            c.alignment = Alignment(horizontal="left" if i < 3 else "center", vertical="center", indent=1 if i < 3 else 0)
            c.border = Border(top=side(GROUP_LINE, "medium") if new_seminar and ti else side(), bottom=side())
            if ti % 2:
                c.fill = fill(BAND_BG)
        ov.row_dimensions[r].height = 20
    o_last = o_first + len(termine) - 1
    tr = o_last + 1
    ov.cell(tr, 2, "Gesamt")
    for i, letter in ((3, "E"), (4, "F"), (5, "G")):
        ov.cell(tr, 2 + i, f"=SUM({letter}{o_first}:{letter}{o_last})")
    for i in range(7):
        c = ov.cell(tr, 2 + i)
        c.font = font(10, True)
        c.fill = fill(TILE_BG)
        c.border = Border(top=side(HEAD_BG, "medium"), bottom=side(HEAD_BG, "medium"))
        c.alignment = Alignment(horizontal="left" if i < 3 else "center", vertical="center", indent=1 if i < 3 else 0)
    ov.row_dimensions[tr].height = 22

    st = f"H{o_first}:H{o_last}"
    for text, bg, fg in (("offen", OPEN_BG, OPEN_FG), ("teilweise", PART_BG, PART_FG), ("komplett", DONE_BG, DONE_FG)):
        ov.conditional_formatting.add(st, FormulaRule(formula=[f'$H{o_first}="{text}"'], fill=fill(bg),
                                                      font=Font(name=FONT, color=fg, bold=True)))

    if hinweise or ignoriert:
        hr = tr + 2
        ov.cell(hr, 2, f"Hinweise – ohne Kontingent für {standort} (nicht in der Liste)").font = font(11, True, HEAD_BG)
        for i, h in enumerate(["Seminar", "Termin", "Bemerkung"]):
            c = ov.cell(hr + 1, 2 + i, h)
            c.font = font(9, True, MUTED)
            c.border = Border(bottom=side(GROUP_LINE))
            c.alignment = Alignment(indent=1)
        for j, h in enumerate(hinweise):
            for i, v in enumerate([h.seminar, h.termin, h.text]):
                c = ov.cell(hr + 2 + j, 2 + i, v)
                c.font = font(9, color=INK)
                c.border = Border(bottom=side())
                c.alignment = Alignment(vertical="center", indent=1)
        if ignoriert:
            ov.cell(hr + 3 + len(hinweise), 2, "Bewusst nicht berücksichtigt: " + ", ".join(ignoriert) + ".")\
                .font = font(9, italic=True, color=MUTED)

    ov.freeze_panes = f"A{o_first}"
    ov.page_setup.orientation = "landscape"
    ov.page_setup.paperSize = ov.PAPERSIZE_A4
    ov.page_setup.fitToWidth, ov.page_setup.fitToHeight = 1, 0
    ov.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    wb.save(out)
    return [ws.title, ov.title]


def main(src, out, standort="FFM"):
    termine, hinweise, ignoriert = read_pptx(src, standort)
    if not termine:
        raise SystemExit(f"Keine Kontingente für {standort} gefunden.")
    years = sorted({t.termine[0][-4:] for t in termine})
    jahr = years[0] if len(years) == 1 else f"{years[0]}–{years[-1]}"
    with tempfile.TemporaryDirectory(prefix="kontingente-") as tmp:
        raw = os.path.join(tmp, "raw.xlsx")
        calc = os.path.join(tmp, "calc.xlsx")
        sheet_names = build(termine, hinweise, ignoriert, standort, jahr, raw)
        shutil.copy(raw, calc)
        libreoffice_recalc(calc, timeout=180)
        wb = load_workbook(calc, data_only=True)
        errors = [f"{n}!{c.coordinate}" for n in sheet_names for row in wb[n].iter_rows() for c in row
                  if isinstance(c.value, str) and c.value in ERRORS]
        if errors:
            raise RuntimeError(f"Formelfehler: {errors[:20]}")
        inject(raw, calc, out, sheet_names)
    print(f"{len(termine)} Termine, {sum(t.plaetze for t in termine)} Plätze {standort}, "
          f"{len(hinweise)} Hinweise -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
